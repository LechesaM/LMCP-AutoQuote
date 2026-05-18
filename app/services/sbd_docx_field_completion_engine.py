from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Tuple
import json
import re
import shutil


ENGINE_VERSION = "SBD_DOCX_FIELD_COMPLETION_ENGINE_V1"

RUNTIME_DIR = Path("runtime/sbd_docx_field_completion")
OUTPUT_DIR = RUNTIME_DIR / "completed"
REPORT_DIR = RUNTIME_DIR / "reports"

DEFAULT_COMPANY_PROFILE = {
    "company_name": "Lechesa Manaba Consulting and Projects (Pty) Ltd",
    "trading_name": "Lechesa Manaba Consulting and Projects (Pty) Ltd",
    "representative_name": "Lechesa Manaba",
    "capacity": "Director",
    "country": "South Africa",
    "supplier_quote_validity_days": "30",
    "bbbee_status_level": "",
    "csd_number": "",
    "tax_compliance_pin": "",
    "vat_number": "",
    "registration_number": "",
    "telephone": "",
    "cellphone": "",
    "email": "",
    "postal_address": "",
    "physical_address": "",
    "signature_text": "Lechesa Manaba",
}


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _safe_str(value: Any) -> str:
    try:
        if value is None:
            return ""
        return str(value).strip()
    except Exception:
        return ""


def _safe_lower(value: Any) -> str:
    return _safe_str(value).lower()


def _slug(value: Any, limit: int = 120) -> str:
    text = _safe_str(value)
    text = re.sub(r"[^A-Za-z0-9]+", "-", text).strip("-")
    return (text[:limit].strip("-") or "sbd-docx")[:limit]


def _ensure_dirs() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    REPORT_DIR.mkdir(parents=True, exist_ok=True)


def _merge_profile(profile: Dict[str, Any] | None) -> Dict[str, str]:
    merged = dict(DEFAULT_COMPANY_PROFILE)
    if isinstance(profile, dict):
        for key, value in profile.items():
            if value not in (None, ""):
                merged[key] = _safe_str(value)
    return {k: _safe_str(v) for k, v in merged.items()}


def _cell_text(cell: Any) -> str:
    return _safe_str(getattr(cell, "text", ""))


def _set_cell_text(cell: Any, value: str) -> None:
    # python-docx cell.text replaces paragraph content safely.
    cell.text = _safe_str(value)


def _looks_empty(text: str) -> bool:
    t = _safe_str(text)
    if not t:
        return True
    t = t.replace("_", "").replace(".", "").replace(":", "").strip()
    return len(t) == 0


def _contains_any(text: str, terms: List[str]) -> bool:
    low = _safe_lower(text)
    return any(term in low for term in terms)


def _field_value_for_label(label: str, profile: Dict[str, str], rfq_context: Dict[str, Any]) -> str:
    low = _safe_lower(label)

    if any(x in low for x in ["name of bidder", "company name", "name of tenderer", "bidder name"]):
        return profile.get("company_name", "")

    if any(x in low for x in ["trading name"]):
        return profile.get("trading_name", profile.get("company_name", ""))

    if any(x in low for x in ["registration number", "company registration"]):
        return profile.get("registration_number", "")

    if any(x in low for x in ["vat registration", "vat number"]):
        return profile.get("vat_number", "")

    if "tax compliance" in low or "tax pin" in low:
        return profile.get("tax_compliance_pin", "")

    if "csd" in low or "central supplier database" in low:
        return profile.get("csd_number", "")

    if "b-bbee" in low or "bbbee" in low or "bbee" in low:
        return profile.get("bbbee_status_level", "")

    if any(x in low for x in ["postal address"]):
        return profile.get("postal_address", "")

    if any(x in low for x in ["street address", "physical address", "business address"]):
        return profile.get("physical_address", "")

    if any(x in low for x in ["telephone number", "tel no", "tel number"]):
        return profile.get("telephone", "")

    if any(x in low for x in ["cellphone", "cell phone", "mobile"]):
        return profile.get("cellphone", "")

    if "e-mail" in low or "email" in low:
        return profile.get("email", "")

    if any(x in low for x in ["duly authorised", "authorised representative", "representative name", "full name"]):
        return profile.get("representative_name", "")

    if "capacity" in low:
        return profile.get("capacity", "")

    if "signature" in low:
        return profile.get("signature_text") or profile.get("representative_name", "")

    if low.strip() == "date" or low.endswith(" date"):
        return datetime.now().strftime("%Y-%m-%d")

    if "bid number" in low or "tender number" in low or "rfq number" in low:
        return _safe_str(rfq_context.get("reference_number") or rfq_context.get("rfq_number"))

    if "bid description" in low or "description of bid" in low:
        return _safe_str(rfq_context.get("title") or rfq_context.get("description"))

    if "validity" in low:
        return profile.get("supplier_quote_validity_days", "30")

    if "country" in low:
        return profile.get("country", "South Africa")

    return ""


def _fill_adjacent_cells(doc: Any, profile: Dict[str, str], rfq_context: Dict[str, Any]) -> List[Dict[str, Any]]:
    changes: List[Dict[str, Any]] = []

    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            cells = row.cells
            for cell_index, cell in enumerate(cells):
                label = _cell_text(cell)
                value = _field_value_for_label(label, profile, rfq_context)
                if not value:
                    continue

                # Fill next empty/right cell first.
                for next_index in range(cell_index + 1, min(cell_index + 3, len(cells))):
                    current = _cell_text(cells[next_index])
                    if _looks_empty(current) or len(current) < 4:
                        _set_cell_text(cells[next_index], value)
                        changes.append({
                            "table": table_index,
                            "row": row_index,
                            "label_cell": cell_index,
                            "target_cell": next_index,
                            "label": label,
                            "value": value,
                            "method": "adjacent_cell",
                        })
                        break

    return changes


def _fill_inline_placeholders(doc: Any, profile: Dict[str, str], rfq_context: Dict[str, Any]) -> List[Dict[str, Any]]:
    changes: List[Dict[str, Any]] = []

    placeholder_map = {
        "{{company_name}}": profile.get("company_name", ""),
        "{{representative_name}}": profile.get("representative_name", ""),
        "{{capacity}}": profile.get("capacity", ""),
        "{{csd_number}}": profile.get("csd_number", ""),
        "{{tax_compliance_pin}}": profile.get("tax_compliance_pin", ""),
        "{{vat_number}}": profile.get("vat_number", ""),
        "{{registration_number}}": profile.get("registration_number", ""),
        "{{email}}": profile.get("email", ""),
        "{{telephone}}": profile.get("telephone", ""),
        "{{cellphone}}": profile.get("cellphone", ""),
        "{{signature}}": profile.get("signature_text") or profile.get("representative_name", ""),
        "{{date}}": datetime.now().strftime("%Y-%m-%d"),
        "{{rfq_number}}": _safe_str(rfq_context.get("reference_number") or rfq_context.get("rfq_number")),
        "{{bid_description}}": _safe_str(rfq_context.get("title") or rfq_context.get("description")),
    }

    def replace_in_paragraph(paragraph: Any, location: str) -> None:
        original = paragraph.text
        updated = original
        for ph, val in placeholder_map.items():
            if ph in updated and val:
                updated = updated.replace(ph, val)
        if updated != original:
            paragraph.text = updated
            changes.append({
                "location": location,
                "old": original,
                "new": updated,
                "method": "inline_placeholder",
            })

    for i, paragraph in enumerate(doc.paragraphs):
        replace_in_paragraph(paragraph, f"paragraph:{i}")

    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                for pi, paragraph in enumerate(cell.paragraphs):
                    replace_in_paragraph(paragraph, f"table:{ti}:row:{ri}:cell:{ci}:paragraph:{pi}")

    return changes


def _extract_text_stats(doc: Any) -> Dict[str, Any]:
    paragraphs = [p.text for p in doc.paragraphs if _safe_str(p.text)]
    table_count = len(doc.tables)
    text_parts = list(paragraphs)
    for table in doc.tables:
        for row in table.rows:
            text_parts.append(" | ".join(_safe_str(c.text) for c in row.cells))
    text = "\n".join(text_parts)
    low = text.lower()
    return {
        "paragraph_count": len(paragraphs),
        "table_count": table_count,
        "text_length": len(text),
        "signals": {
            "contains_sbd3": "sbd 3" in low or "pricing schedule" in low,
            "contains_sbd6": "sbd 6" in low or "preference points" in low,
            "contains_name_of_bidder": "name of bidder" in low,
            "contains_csd": "csd" in low or "central supplier database" in low,
            "contains_signature": "signature" in low,
            "contains_bbbee": "b-bbee" in low or "bbbee" in low or "bbee" in low,
        },
    }


def complete_sbd_docx(
    sbd_docx_path: str | Path,
    company_profile: Dict[str, Any] | None = None,
    rfq_context: Dict[str, Any] | None = None,
) -> Dict[str, Any]:
    _ensure_dirs()

    try:
        from docx import Document
    except Exception as exc:
        return {
            "status": "failed",
            "engine_version": ENGINE_VERSION,
            "error": f"python-docx unavailable: {exc}",
            "output_docx_path": "",
            "report_path": "",
        }

    source = Path(_safe_str(sbd_docx_path))
    profile = _merge_profile(company_profile)
    context = rfq_context if isinstance(rfq_context, dict) else {}

    if not source.exists():
        return {
            "status": "failed",
            "engine_version": ENGINE_VERSION,
            "error": "sbd_docx_path_not_found",
            "sbd_docx_path": str(source),
            "output_docx_path": "",
            "report_path": "",
        }

    output_name = f"{_slug(source.stem)}__completed.docx"
    output_path = OUTPUT_DIR / output_name
    report_path = REPORT_DIR / f"{_slug(source.stem)}__completion_report.json"

    shutil.copy2(source, output_path)
    doc = Document(str(output_path))

    before_stats = _extract_text_stats(doc)

    changes: List[Dict[str, Any]] = []
    changes.extend(_fill_inline_placeholders(doc, profile, context))
    changes.extend(_fill_adjacent_cells(doc, profile, context))

    doc.save(str(output_path))

    # Reopen for after stats
    doc_after = Document(str(output_path))
    after_stats = _extract_text_stats(doc_after)

    missing_profile_fields = [
        key for key in (
            "registration_number", "vat_number", "tax_compliance_pin", "csd_number",
            "bbbee_status_level", "telephone", "cellphone", "email",
            "postal_address", "physical_address",
        )
        if not profile.get(key)
    ]

    result = {
        "status": "ok",
        "engine_version": ENGINE_VERSION,
        "completed_at": _now_iso(),
        "sbd_docx_path": str(source),
        "output_docx_path": str(output_path),
        "profile_company_name": profile.get("company_name"),
        "rfq_reference": _safe_str(context.get("reference_number") or context.get("rfq_number")),
        "changes_count": len(changes),
        "changes": changes[:500],
        "before_stats": before_stats,
        "after_stats": after_stats,
        "missing_profile_fields": missing_profile_fields,
        "manual_review_required": True,
        "final_submit_allowed": False,
        "notes": [
            "Completed DOCX copy created; original SBD was not overwritten.",
            "Blank company profile fields require manual review/completion.",
            "Signature image/handwriting overlay is handled by the PDF/SBD intelligence stack, not this DOCX writer.",
        ],
    }

    report_path.write_text(json.dumps(result, indent=2, default=str))
    result["report_path"] = str(report_path)
    return result


def get_status() -> Dict[str, Any]:
    return {
        "status": "ok",
        "engine_version": ENGINE_VERSION,
        "runtime_dir": str(RUNTIME_DIR),
        "output_dir": str(OUTPUT_DIR),
        "report_dir": str(REPORT_DIR),
        "capabilities": [
            "complete_sbd_docx_copy",
            "fill_adjacent_cells_from_labels",
            "fill_inline_placeholders",
            "preserve_original_docx",
            "write_completion_report",
            "manual_review_required",
        ],
    }
