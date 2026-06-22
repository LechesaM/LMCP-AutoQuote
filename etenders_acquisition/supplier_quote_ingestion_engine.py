#!/usr/bin/env python3
import csv
import json
import re
from pathlib import Path
from datetime import datetime, timezone
from collections import Counter, defaultdict

from openpyxl import load_workbook

RUNTIME_DIR = Path("/Users/cash/Documents/runtime")

RFQ_PACK_INDEX_FILE = RUNTIME_DIR / "rfq_packs" / "rfq_pack_index.json"

SUPPLIER_QUOTES_DIR = RUNTIME_DIR / "supplier_quotes"
OUT_DIR = RUNTIME_DIR / "supplier_quote_ingestion"

QUOTE_ITEMS_FILE = OUT_DIR / "supplier_quote_items.json"
COMPARISON_FILE = OUT_DIR / "supplier_quote_comparison_matrix.json"
SUMMARY_FILE = OUT_DIR / "supplier_quote_ingestion_summary.json"

EXPECTED_COLUMNS = {
    "line_no": ["line_no", "line", "item", "item_no", "no"],
    "description": ["description", "item_description", "desc"],
    "quantity": ["quantity", "qty"],
    "unit": ["unit", "uom"],
    "unit_price": ["unit_price", "rate", "price", "quoted_price", "supplier_price"],
    "total_price": ["total_price", "total", "amount", "line_total"],
    "availability": ["availability", "available", "stock"],
    "lead_time": ["lead_time", "delivery", "delivery_time"],
    "notes": ["notes", "comments", "remark", "remarks"],
}


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def clean(value):
    value = "" if value is None else str(value)
    value = value.replace("\n", " ").replace("\r", " ")
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def norm_key(value):
    return re.sub(r"[^a-z0-9]+", "_", clean(value).lower()).strip("_")


def numeric(value):
    if value is None:
        return None

    if isinstance(value, (int, float)):
        return float(value)

    value = clean(value)
    value = value.replace(",", "")
    value = value.replace("R", "").replace("r", "")

    value = re.sub(r"[^\d.\-]", "", value)

    if value in ["", ".", "-", "-."]:
        return None

    try:
        return float(value)
    except Exception:
        return None


def load_json(path, default=None):
    if default is None:
        default = {}

    if not path.exists():
        return default

    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)

    with path.open("w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def supplier_from_path(path):
    return path.parent.name.replace("_", " ")


def find_quote_files():
    SUPPLIER_QUOTES_DIR.mkdir(parents=True, exist_ok=True)

    files = []

    for ext in ["*.csv", "*.xlsx", "*.xls"]:
        files.extend(SUPPLIER_QUOTES_DIR.rglob(ext))

    return sorted(files)


def map_columns(headers):
    normalized_headers = [norm_key(h) for h in headers]
    mapping = {}

    for target, aliases in EXPECTED_COLUMNS.items():
        for alias in aliases:
            alias_norm = norm_key(alias)
            if alias_norm in normalized_headers:
                mapping[target] = normalized_headers.index(alias_norm)
                break

    return mapping


def parse_csv(path):
    rows = []

    with path.open("r", encoding="utf-8-sig", errors="ignore") as f:
        reader = csv.reader(f)
        raw = list(reader)

    if not raw:
        return []

    headers = raw[0]
    mapping = map_columns(headers)

    for row in raw[1:]:
        rows.append(row_to_quote_item(path, row, mapping))

    return [r for r in rows if r]


def parse_excel(path):
    wb = load_workbook(path, data_only=True, read_only=True)

    all_items = []

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            continue

        header_idx = None
        mapping = {}

        for i, row in enumerate(rows[:20]):
            values = [clean(v) for v in row]
            candidate_mapping = map_columns(values)

            if "description" in candidate_mapping and (
                "unit_price" in candidate_mapping or "total_price" in candidate_mapping
            ):
                header_idx = i
                mapping = candidate_mapping
                break

        if header_idx is None:
            continue

        for row in rows[header_idx + 1:]:
            item = row_to_quote_item(path, list(row), mapping)
            if item:
                item["sheet"] = ws.title
                all_items.append(item)

    return all_items


def row_value(row, mapping, key):
    idx = mapping.get(key)

    if idx is None or idx >= len(row):
        return ""

    return row[idx]


def row_to_quote_item(path, row, mapping):
    description = clean(row_value(row, mapping, "description"))

    if not description or len(description) < 3:
        return None

    quantity = numeric(row_value(row, mapping, "quantity"))
    unit_price = numeric(row_value(row, mapping, "unit_price"))
    total_price = numeric(row_value(row, mapping, "total_price"))

    if total_price is None and quantity is not None and unit_price is not None:
        total_price = round(quantity * unit_price, 2)

    return {
        "supplier": supplier_from_path(path),
        "source_file": str(path),
        "line_no": clean(row_value(row, mapping, "line_no")),
        "description": description,
        "quantity": quantity,
        "unit": clean(row_value(row, mapping, "unit")),
        "unit_price": unit_price,
        "total_price": total_price,
        "availability": clean(row_value(row, mapping, "availability")),
        "lead_time": clean(row_value(row, mapping, "lead_time")),
        "notes": clean(row_value(row, mapping, "notes")),
        "ingested_at": now_iso(),
    }


def description_key(description):
    value = clean(description).lower()
    value = re.sub(r"vendors are responsible.*", "", value)
    value = re.sub(r"[^a-z0-9]+", " ", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip()[:180]


def build_comparison(items):
    grouped = defaultdict(list)

    for item in items:
        key = description_key(item.get("description"))
        grouped[key].append(item)

    matrix = []

    for key, group in grouped.items():
        priced = [
            g for g in group
            if g.get("unit_price") is not None or g.get("total_price") is not None
        ]

        if priced:
            best = min(
                priced,
                key=lambda x: (
                    x.get("total_price")
                    if x.get("total_price") is not None
                    else x.get("unit_price")
                ),
            )
        else:
            best = None

        suppliers = sorted(set(g.get("supplier") for g in group))

        matrix.append({
            "description_key": key,
            "description": group[0].get("description"),
            "quote_count": len(group),
            "suppliers": suppliers,
            "best_supplier": best.get("supplier") if best else None,
            "best_unit_price": best.get("unit_price") if best else None,
            "best_total_price": best.get("total_price") if best else None,
            "quotes": group,
        })

    return sorted(
        matrix,
        key=lambda x: (
            x["best_total_price"] is None,
            x["best_total_price"] or 0,
        ),
    )


def main():
    quote_files = find_quote_files()

    all_items = []
    failed = []

    for path in quote_files:
        print(f"Ingesting: {path}")

        try:
            if path.suffix.lower() == ".csv":
                items = parse_csv(path)
            else:
                items = parse_excel(path)

            all_items.extend(items)

        except Exception as e:
            failed.append({
                "file": str(path),
                "error": str(e),
            })

    comparison = build_comparison(all_items)

    supplier_counts = Counter(i.get("supplier") for i in all_items)
    priced_items = [
        i for i in all_items
        if i.get("unit_price") is not None or i.get("total_price") is not None
    ]

    summary = {
        "generated_at": now_iso(),
        "quote_files_found": len(quote_files),
        "quote_files_failed": len(failed),
        "items_ingested": len(all_items),
        "priced_items": len(priced_items),
        "comparison_lines": len(comparison),
        "supplier_counts": dict(supplier_counts),
        "failed_files": failed,
        "input_directory": str(SUPPLIER_QUOTES_DIR),
    }

    output = {
        "generated_at": now_iso(),
        "items": all_items,
    }

    write_json(QUOTE_ITEMS_FILE, output)
    write_json(COMPARISON_FILE, {
        "generated_at": now_iso(),
        "comparison": comparison,
    })
    write_json(SUMMARY_FILE, summary)

    print(f"\nSupplier quote ingestion summary: {SUMMARY_FILE}")
    print(f"Quote files found: {summary['quote_files_found']}")
    print(f"Items ingested: {summary['items_ingested']}")
    print(f"Priced items: {summary['priced_items']}")
    print(f"Comparison lines: {summary['comparison_lines']}")
    print(f"Failed files: {summary['quote_files_failed']}")
    print(f"Input folder: {SUPPLIER_QUOTES_DIR}")


if __name__ == "__main__":
    main()
