#!/usr/bin/env python3
import json
import re
from pathlib import Path
from datetime import datetime, timezone
from collections import Counter

from openpyxl import load_workbook

RUNTIME_DIR = Path("/Users/cash/Documents/runtime")

INVENTORY_FILE = RUNTIME_DIR / "boq_intelligence" / "document_type_inventory.json"
OUT_DIR = RUNTIME_DIR / "boq_intelligence"
OUTPUT_FILE = OUT_DIR / "excel_boq_parse_results.json"

MAX_ROWS_PER_SHEET = 5000

HEADER_HINTS = [
    "description", "item", "material", "goods", "service", "scope",
    "unit", "uom", "qty", "quantity", "rate", "price", "amount",
    "total", "boq", "bill", "schedule"
]

DESCRIPTION_HINTS = [
    "description", "item description", "material", "goods", "service",
    "scope", "specification", "product", "commodity"
]

QTY_HINTS = ["qty", "quantity", "quantities", "no.", "number"]
UNIT_HINTS = ["unit", "uom", "measure", "measurement"]
RATE_HINTS = ["rate", "unit price", "price", "amount", "total", "value"]


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def load_json(path, default=None):
    if default is None:
        default = {}
    if not path.exists():
        return default
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def clean(value):
    value = "" if value is None else str(value)
    value = value.replace("\n", " ").replace("\r", " ")
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def clean_lower(value):
    return clean(value).lower()


def is_empty_row(values):
    return all(clean(v) == "" for v in values)


def row_text(values):
    return " ".join(clean(v) for v in values if clean(v))


def looks_like_header(values):
    text = clean_lower(row_text(values))
    if not text:
        return False

    score = 0
    for h in HEADER_HINTS:
        if h in text:
            score += 1

    return score >= 2


def find_header_row(rows):
    best = None

    for idx, values in enumerate(rows[:80]):
        if is_empty_row(values):
            continue

        text = clean_lower(row_text(values))
        score = 0

        for h in HEADER_HINTS:
            if h in text:
                score += 1

        non_empty = sum(1 for v in values if clean(v))
        if non_empty >= 3:
            score += 1

        if score >= 3 and (best is None or score > best["score"]):
            best = {
                "row_index": idx,
                "score": score,
                "values": [clean(v) for v in values],
            }

    return best


def find_column(headers, hints):
    best_idx = None
    best_score = 0

    for i, h in enumerate(headers):
        hl = clean_lower(h)
        score = 0

        for hint in hints:
            if hint == hl:
                score += 5
            elif hint in hl:
                score += 3

        if score > best_score:
            best_score = score
            best_idx = i

    return best_idx, best_score


def numeric_value(value):
    if value is None:
        return None

    if isinstance(value, (int, float)):
        return float(value)

    s = clean(value)
    if not s:
        return None

    s = s.replace(",", "")
    s = s.replace("R", "").replace("r", "")
    s = re.sub(r"[^\d.\-]", "", s)

    if s in ["", ".", "-", "-."]:
        return None

    try:
        return float(s)
    except Exception:
        return None


def extract_rows_from_sheet(ws):
    rows = []
    max_row = min(ws.max_row or 0, MAX_ROWS_PER_SHEET)
    max_col = ws.max_column or 0

    for r in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True):
        values = list(r)
        if not is_empty_row(values):
            rows.append(values)

    return rows


def parse_sheet(ws):
    rows = extract_rows_from_sheet(ws)

    if not rows:
        return {
            "sheet_name": ws.title,
            "row_count": 0,
            "header_found": False,
            "items": [],
            "warnings": ["empty_sheet"],
        }

    header = find_header_row(rows)

    if not header:
        return {
            "sheet_name": ws.title,
            "row_count": len(rows),
            "header_found": False,
            "items": [],
            "warnings": ["no_header_detected"],
            "sample_rows": [[clean(v) for v in row[:12]] for row in rows[:10]],
        }

    header_idx = header["row_index"]
    headers = header["values"]

    desc_col, desc_score = find_column(headers, DESCRIPTION_HINTS)
    qty_col, qty_score = find_column(headers, QTY_HINTS)
    unit_col, unit_score = find_column(headers, UNIT_HINTS)
    rate_col, rate_score = find_column(headers, RATE_HINTS)

    items = []

    for absolute_idx, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        values = [clean(v) for v in row]

        if is_empty_row(values):
            continue

        description = values[desc_col] if desc_col is not None and desc_col < len(values) else ""
        qty_raw = values[qty_col] if qty_col is not None and qty_col < len(values) else ""
        unit = values[unit_col] if unit_col is not None and unit_col < len(values) else ""
        rate_raw = values[rate_col] if rate_col is not None and rate_col < len(values) else ""

        qty = numeric_value(qty_raw)
        rate = numeric_value(rate_raw)

        text = clean(row_text(values))

        if len(text) < 5:
            continue

        if not description and len(values) > 0:
            description = max(values, key=lambda x: len(clean(x)))

        if not description or len(description) < 3:
            continue

        item = {
            "source_row_number": absolute_idx,
            "description": description,
            "quantity": qty,
            "quantity_raw": qty_raw,
            "unit": unit,
            "rate": rate,
            "rate_raw": rate_raw,
            "row_values": values[:30],
        }

        # Keep likely item rows only
        if qty is not None or unit or len(description) > 10:
            items.append(item)

    return {
        "sheet_name": ws.title,
        "row_count": len(rows),
        "header_found": True,
        "header_row_number": header_idx + 1,
        "header_values": headers,
        "column_map": {
            "description_col": desc_col,
            "quantity_col": qty_col,
            "unit_col": unit_col,
            "rate_col": rate_col,
            "scores": {
                "description": desc_score,
                "quantity": qty_score,
                "unit": unit_score,
                "rate": rate_score,
            },
        },
        "items_extracted": len(items),
        "items": items,
        "warnings": [],
    }


def parse_workbook(path):
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as e:
        return {
            "path": str(path),
            "filename": path.name,
            "success": False,
            "error": str(e),
            "sheets": [],
            "total_items": 0,
        }

    sheets = []
    total_items = 0

    for ws in wb.worksheets:
        parsed = parse_sheet(ws)
        sheets.append(parsed)
        total_items += parsed.get("items_extracted", 0)

    return {
        "path": str(path),
        "filename": path.name,
        "success": True,
        "sheet_count": len(sheets),
        "total_items": total_items,
        "sheets": sheets,
    }


def main():
    inventory = load_json(INVENTORY_FILE, default={})

    candidates = inventory.get("price_candidates", [])

    excel_candidates = [
        Path(c["path"])
        for c in candidates
        if c.get("extension") in [".xlsx", ".xls", ".csv"]
    ]

    results = []

    for path in excel_candidates:
        print(f"Parsing: {path.name}")
        results.append(parse_workbook(path))

    summary = {
        "generated_at": now_iso(),
        "excel_candidates": len(excel_candidates),
        "workbooks_parsed": sum(1 for r in results if r.get("success")),
        "workbooks_failed": sum(1 for r in results if not r.get("success")),
        "total_items_extracted": sum(r.get("total_items", 0) for r in results),
        "results": results,
    }

    write_json(OUTPUT_FILE, summary)

    print(f"\nExcel BOQ parse results written: {OUTPUT_FILE}")
    print(f"Excel candidates: {summary['excel_candidates']}")
    print(f"Parsed: {summary['workbooks_parsed']}")
    print(f"Failed: {summary['workbooks_failed']}")
    print(f"Items extracted: {summary['total_items_extracted']}")


if __name__ == "__main__":
    main()
