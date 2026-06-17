from __future__ import annotations

import json
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from app.services import live_rfq_store
from app.services import rfq_zip_content_extraction_engine as zip_engine
from app.services.rfq_docx_main_document_intelligence import analyse_docx_main_document
from app.services.tender_harvester import _lmcp_apply_docx_verified_quantity_gate


def _write_docx(path: Path, paragraphs: list[str], tables: list[list[list[str]]] | None = None) -> Path:
    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    for table_rows in tables or []:
        table = doc.add_table(rows=0, cols=len(table_rows[0]))
        for row in table_rows:
            cells = table.add_row().cells
            for idx, value in enumerate(row):
                cells[idx].text = value
    doc.save(path)
    return path


def _write_xlsx(path: Path, rows: list[list[str]]) -> Path:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def _write_csv(path: Path, rows: list[list[str]]) -> Path:
    path.write_text("\n".join(",".join(row) for row in rows), encoding="utf-8")
    return path


def _write_fake_pdf(path: Path, text: str) -> Path:
    path.write_bytes(f"%PDF-1.4\n{text}\n%%EOF".encode("utf-8"))
    return path


def _build_zip_bundle(tmp_path: Path, files: list[Path], name: str = "buyer-pack.zip") -> Path:
    zip_path = tmp_path / name
    with zipfile.ZipFile(zip_path, "w") as zf:
        for file_path in files:
            zf.write(file_path, arcname=file_path.name)
    return zip_path


def _patch_zip_runtime(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setattr(zip_engine, "RUNTIME_DIR", tmp_path / "runtime")
    monkeypatch.setattr(zip_engine, "EXTRACTED_DIR", tmp_path / "runtime" / "extracted")
    monkeypatch.setattr(zip_engine, "REPORT_DIR", tmp_path / "runtime" / "reports")


def test_main_document_path_missing_does_not_crash(monkeypatch) -> None:
    monkeypatch.setattr(
        "app.services.tender_harvester.acquire_rfq_documents",
        lambda payload: {"status": "ok", "confidence": 0.9},
    )
    monkeypatch.setattr(
        "app.services.tender_harvester.extract_zip_contents",
        lambda acquisition: {
            "status": "ok",
            "confidence": 0.8,
            "main_document_path": "",
            "extracted_files": [],
            "sbd_document_paths": [],
            "boq_candidate_paths": [],
            "pricing_schedule_paths": [],
            "specification_paths": [],
        },
    )
    monkeypatch.setattr(
        "app.services.tender_harvester.analyse_docx_main_document",
        lambda path: {"status": "ok"},
    )

    item = _lmcp_apply_docx_verified_quantity_gate(
        {
            "eligible": True,
            "document_url": "https://example.com/rfq",
            "buyer_rfq_number": "RFQ-MISSING-MAIN",
            "title": "Missing main document pack",
        }
    )

    assert item["buyer_pack_downloaded"] is False
    assert item["docx_main_document_intelligence_result"]["status"] == "skipped"
    assert item["docx_main_document_intelligence_result"]["reason"] == "main_document_path_not_found"


def test_zip_buyer_pack_inventory_is_summarized(tmp_path: Path, monkeypatch) -> None:
    _patch_zip_runtime(monkeypatch, tmp_path)
    docx_path = _write_docx(tmp_path / "1. Main.docx", ["Pricing/Billing Model"])
    sbd_path = _write_docx(tmp_path / "2. SBD.docx", ["SBD 3.2", "PRICING SCHEDULE", "Tax compliance"])
    xlsx_path = _write_xlsx(tmp_path / "Pricing Data.xlsx", [["Item", "Description", "Qty", "Rate", "Amount"], ["1", "Pump", "2", "10", "20"]])
    csv_path = _write_csv(tmp_path / "rates.csv", [["item", "quantity", "rate", "amount"], ["Valve", "3", "12", "36"]])
    pdf_path = _write_fake_pdf(tmp_path / "Scope.pdf", "Scope and pricing")
    zip_path = _build_zip_bundle(tmp_path, [docx_path, sbd_path, xlsx_path, csv_path, pdf_path])

    result = zip_engine.extract_zip_contents(
        {
            "title": "Inventory Pack",
            "buyer_rfq_number": "RFQ-INVENTORY",
            "downloaded_files": [{"path": str(zip_path), "extension": ".zip"}],
        }
    )

    assert result["artifact_count"] == 5
    assert result["pdf_count"] == 1
    assert result["docx_count"] == 2
    assert result["xlsx_count"] == 1
    assert result["csv_count"] == 1
    assert result["zip_count"] == 1
    assert result["extracted_file_count"] == 5
    assert result["main_document_path"].endswith("1. Main.docx")
    assert "docx" in result["detected_document_types"]
    assert len(result["document_inventory_paths_limited"]) == 5


def test_boq_detected_by_filename(tmp_path: Path, monkeypatch) -> None:
    _patch_zip_runtime(monkeypatch, tmp_path)
    boq_sheet = _write_xlsx(tmp_path / "BOQ.xlsx", [["Item", "Description", "Qty", "Amount"], ["1", "Cable", "4", "100"]])
    zip_path = _build_zip_bundle(tmp_path, [boq_sheet], "boq-filename.zip")

    result = zip_engine.extract_zip_contents(
        {"title": "BOQ filename", "downloaded_files": [{"path": str(zip_path), "extension": ".zip"}]}
    )

    assert result["boq_detected"] is True
    assert result["boq_detection_confidence"] >= 0.55


def test_boq_detected_by_content_phrase(tmp_path: Path, monkeypatch) -> None:
    _patch_zip_runtime(monkeypatch, tmp_path)
    neutral_doc = _write_docx(
        tmp_path / "neutral.docx",
        ["This document contains the Bill of Quantities for the required works."],
    )
    zip_path = _build_zip_bundle(tmp_path, [neutral_doc], "boq-content.zip")

    result = zip_engine.extract_zip_contents(
        {"title": "BOQ content", "downloaded_files": [{"path": str(zip_path), "extension": ".zip"}]}
    )

    assert result["boq_detected"] is True
    assert "boq_terms" in result["boq_detection_reason"]


def test_pricing_schedule_detected_by_xlsx(tmp_path: Path, monkeypatch) -> None:
    _patch_zip_runtime(monkeypatch, tmp_path)
    pricing_sheet = _write_xlsx(
        tmp_path / "schedule.xlsx",
        [["Item", "Description", "Quantity", "Rate", "Amount"], ["1", "Steel", "5", "20", "100"]],
    )
    zip_path = _build_zip_bundle(tmp_path, [pricing_sheet], "pricing-xlsx.zip")

    result = zip_engine.extract_zip_contents(
        {"title": "Pricing xlsx", "downloaded_files": [{"path": str(zip_path), "extension": ".zip"}]}
    )

    assert result["pricing_schedule_detected"] is True
    assert result["pricing_schedule_detection_confidence"] >= 0.55


def test_pricing_schedule_detected_by_sbd_pricing_data_phrase(tmp_path: Path) -> None:
    sbd_doc = _write_docx(
        tmp_path / "sbd-pricing.docx",
        ["SBD 3.1", "Pricing Data", "Form of Offer"],
        tables=[[["Item", "Description", "Rate", "Amount"], ["1", "Service", "100", "100"]]],
    )

    result = analyse_docx_main_document(sbd_doc)

    assert result["pricing_schedule_detected"] is True
    assert result["pricing_schedule_detection_confidence"] >= 0.55


def test_returnables_detection_remains_true_for_sbd_bundle(tmp_path: Path, monkeypatch) -> None:
    _patch_zip_runtime(monkeypatch, tmp_path)
    sbd_doc = _write_docx(
        tmp_path / "2. SBD.docx",
        ["Returnable documents checklist", "Tax compliance", "CSD", "B-BBEE"],
    )
    zip_path = _build_zip_bundle(tmp_path, [sbd_doc], "returnables.zip")

    result = zip_engine.extract_zip_contents(
        {"title": "Returnables bundle", "downloaded_files": [{"path": str(zip_path), "extension": ".zip"}]}
    )

    assert result["returnables_detected"] is True
    assert result["returnables_detection_confidence"] >= 0.45


def test_quote_pack_readiness_score_increases_when_boq_and_pricing_detected() -> None:
    summary = live_rfq_store.summarize_rfq_document_intelligence(
        {
            "buyer_pack_downloaded": True,
            "zip_content_extraction_result": {
                "boq_detected": True,
                "pricing_schedule_detected": True,
                "returnables_detected": True,
                "boq_detection_confidence": 0.9,
                "pricing_schedule_detection_confidence": 0.92,
                "returnables_detection_confidence": 0.75,
            },
        }
    )

    assert summary["boq_detected"] is True
    assert summary["pricing_schedule_detected"] is True
    assert summary["returnables_detected"] is True
    assert summary["quote_pack_readiness_score"] >= 80


def test_quote_pack_artifact_evidence_is_summarized(tmp_path: Path, monkeypatch) -> None:
    quote_pack_dir = tmp_path / "quote-pack"
    quote_pack_dir.mkdir()
    artifact = quote_pack_dir / "formal_quotation_v44.html"
    artifact.write_text("<html>quote</html>", encoding="utf-8")

    monkeypatch.setattr(live_rfq_store, "LIVE_RFQ_STORE_PATH", tmp_path / "live_rfqs.json")

    summary = live_rfq_store.summarize_rfq_document_intelligence(
        {
            "buyer_pack_downloaded": True,
            "boq_detected": True,
            "pricing_schedule_detected": True,
            "returnables_detected": True,
            "quote_pack_path": str(quote_pack_dir),
            "quote_pack_generated": True,
        }
    )

    assert summary["quote_pack_generated"] is True
    assert summary["quote_pack_artifact_exists"] is True
    assert summary["quote_pack_artifact_size_bytes"] >= artifact.stat().st_size
    assert summary["quote_pack_path"] == str(quote_pack_dir)
