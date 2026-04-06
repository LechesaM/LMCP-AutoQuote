import json
import os
import re
import subprocess
import tempfile
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from pypdf import PdfReader, PdfWriter
from pypdf.generic import BooleanObject, NameObject

from docx import Document
from docx.shared import Inches

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage

from reportlab.lib.colors import black, white
from reportlab.pdfgen import canvas


DEFAULT_OUTPUT_DIR = "runtime/generated_forms"
DEFAULT_PROFILE_DIR = "app/data/form_profiles"


@dataclass
class PDFFieldCoordinate:
    page: int
    x: float
    y: float
    width: Optional[float] = None
    height: Optional[float] = None
    font_size: int = 10
    multiline: bool = False
    whiteout: bool = True
    whiteout_padding: float = 2.0
    max_lines: Optional[int] = None
    align: str = "left"


@dataclass
class SignatureCoordinate:
    page: int
    x: float
    y: float
    width: float
    height: float


@dataclass
class AnchorRule:
    field_name: str
    anchor_text: str
    page: Optional[int] = None
    x_offset: float = 0.0
    y_offset: float = 0.0
    width: Optional[float] = None
    height: Optional[float] = None
    font_size: int = 10
    multiline: bool = False
    whiteout: bool = True
    whiteout_padding: float = 2.0
    match_mode: str = "contains"
    max_lines: Optional[int] = None
    align: str = "left"


@dataclass
class FormProfile:
    name: str
    template_type: str
    field_aliases: Dict[str, List[str]] = field(default_factory=dict)
    pdf_coordinates: Dict[str, PDFFieldCoordinate] = field(default_factory=dict)
    anchor_rules: List[AnchorRule] = field(default_factory=list)
    signature_coordinates: Optional[SignatureCoordinate] = None
    docx_signature_placeholder: Optional[str] = None
    xlsx_signature_anchor: Optional[str] = None


class UniversalFormFillerError(Exception):
    pass


class UniversalFormFiller:
    """
    Hybrid universal form filler.

    PDF priority:
    1. Editable PDF form fields
    2. Anchor-based overlay placement
    3. Coordinate-based overlay placement

    DOCX / XLSX:
    - placeholder replacement
    - optional signature insertion
    - optional PDF conversion through LibreOffice
    """

    def __init__(
        self,
        output_dir: str = DEFAULT_OUTPUT_DIR,
        profile_dir: str = DEFAULT_PROFILE_DIR,
        libreoffice_binary: str = "soffice",
    ) -> None:
        self.output_dir = Path(output_dir)
        self.profile_dir = Path(profile_dir)
        self.libreoffice_binary = libreoffice_binary
        self.output_dir.mkdir(parents=True, exist_ok=True)

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def fill_form(
        self,
        input_path: str,
        data: Dict[str, Any],
        output_basename: Optional[str] = None,
        profile_name: Optional[str] = None,
        signature_path: Optional[str] = None,
        convert_to_pdf: bool = True,
    ) -> Dict[str, Any]:
        input_file = Path(input_path)
        if not input_file.exists():
            raise UniversalFormFillerError(f"Input file not found: {input_path}")

        ext = input_file.suffix.lower()
        profile = self._load_profile(profile_name) if profile_name else None

        safe_basename = output_basename or f"{input_file.stem}_filled_{uuid.uuid4().hex[:8]}"
        safe_basename = self._sanitize_filename(safe_basename)

        if ext == ".pdf":
            filled_path = self.output_dir / f"{safe_basename}.pdf"
            result = self._fill_pdf(
                input_path=str(input_file),
                output_path=str(filled_path),
                data=data,
                profile=profile,
                signature_path=signature_path,
            )
            return {
                "status": "success",
                "input_path": str(input_file),
                "output_path": str(filled_path),
                "output_pdf_path": str(filled_path),
                "details": result,
            }

        if ext == ".docx":
            filled_docx_path = self.output_dir / f"{safe_basename}.docx"
            self._fill_docx(
                input_path=str(input_file),
                output_path=str(filled_docx_path),
                data=data,
                profile=profile,
                signature_path=signature_path,
            )

            output_pdf_path = None
            if convert_to_pdf:
                output_pdf_path = self._convert_office_to_pdf(str(filled_docx_path))

            return {
                "status": "success",
                "input_path": str(input_file),
                "output_path": str(filled_docx_path),
                "output_pdf_path": output_pdf_path,
                "details": {
                    "type": "docx",
                    "converted_to_pdf": bool(output_pdf_path),
                },
            }

        if ext in {".xlsx", ".xlsm"}:
            filled_xlsx_path = self.output_dir / f"{safe_basename}.xlsx"
            self._fill_xlsx(
                input_path=str(input_file),
                output_path=str(filled_xlsx_path),
                data=data,
                profile=profile,
                signature_path=signature_path,
            )

            output_pdf_path = None
            if convert_to_pdf:
                output_pdf_path = self._convert_office_to_pdf(str(filled_xlsx_path))

            return {
                "status": "success",
                "input_path": str(input_file),
                "output_path": str(filled_xlsx_path),
                "output_pdf_path": output_pdf_path,
                "details": {
                    "type": "xlsx",
                    "converted_to_pdf": bool(output_pdf_path),
                },
            }

        raise UniversalFormFillerError(
            f"Unsupported file type: {ext}. Supported: .pdf, .docx, .xlsx, .xlsm"
        )

    # ------------------------------------------------------------------
    # Profile loading
    # ------------------------------------------------------------------

    def _load_profile(self, profile_name: str) -> FormProfile:
        profile_path = self.profile_dir / profile_name
        if not profile_path.exists():
            raise UniversalFormFillerError(f"Profile not found: {profile_path}")

        with open(profile_path, "r", encoding="utf-8") as f:
            raw = json.load(f)

        pdf_coordinates: Dict[str, PDFFieldCoordinate] = {}
        for key, value in raw.get("pdf_coordinates", {}).items():
            pdf_coordinates[key] = PDFFieldCoordinate(
                page=int(value["page"]),
                x=float(value["x"]),
                y=float(value["y"]),
                width=float(value["width"]) if value.get("width") is not None else None,
                height=float(value["height"]) if value.get("height") is not None else None,
                font_size=int(value.get("font_size", 10)),
                multiline=bool(value.get("multiline", False)),
                whiteout=bool(value.get("whiteout", True)),
                whiteout_padding=float(value.get("whiteout_padding", 2.0)),
                max_lines=int(value["max_lines"]) if value.get("max_lines") is not None else None,
                align=str(value.get("align", "left")),
            )

        anchor_rules: List[AnchorRule] = []
        for item in raw.get("anchor_rules", []):
            anchor_rules.append(
                AnchorRule(
                    field_name=item["field_name"],
                    anchor_text=item["anchor_text"],
                    page=int(item["page"]) if item.get("page") is not None else None,
                    x_offset=float(item.get("x_offset", 0.0)),
                    y_offset=float(item.get("y_offset", 0.0)),
                    width=float(item["width"]) if item.get("width") is not None else None,
                    height=float(item["height"]) if item.get("height") is not None else None,
                    font_size=int(item.get("font_size", 10)),
                    multiline=bool(item.get("multiline", False)),
                    whiteout=bool(item.get("whiteout", True)),
                    whiteout_padding=float(item.get("whiteout_padding", 2.0)),
                    match_mode=str(item.get("match_mode", "contains")),
                    max_lines=int(item["max_lines"]) if item.get("max_lines") is not None else None,
                    align=str(item.get("align", "left")),
                )
            )

        signature_coordinates = None
        sig = raw.get("signature_coordinates")
        if sig:
            signature_coordinates = SignatureCoordinate(
                page=int(sig["page"]),
                x=float(sig["x"]),
                y=float(sig["y"]),
                width=float(sig["width"]),
                height=float(sig["height"]),
            )

        return FormProfile(
            name=raw["name"],
            template_type=raw.get("template_type", "generic"),
            field_aliases=raw.get("field_aliases", {}),
            pdf_coordinates=pdf_coordinates,
            anchor_rules=anchor_rules,
            signature_coordinates=signature_coordinates,
            docx_signature_placeholder=raw.get("docx_signature_placeholder"),
            xlsx_signature_anchor=raw.get("xlsx_signature_anchor"),
        )

    # ------------------------------------------------------------------
    # PDF filling
    # ------------------------------------------------------------------

    def _fill_pdf(
        self,
        input_path: str,
        output_path: str,
        data: Dict[str, Any],
        profile: Optional[FormProfile] = None,
        signature_path: Optional[str] = None,
    ) -> Dict[str, Any]:
        reader = PdfReader(input_path)
        fields = self._extract_pdf_fields(reader)

        if fields:
            return self._fill_editable_pdf(
                reader=reader,
                output_path=output_path,
                fields=fields,
                data=data,
                profile=profile,
                signature_path=signature_path,
            )

        if profile and profile.anchor_rules:
            try:
                return self._fill_anchor_pdf(
                    reader=reader,
                    output_path=output_path,
                    data=data,
                    profile=profile,
                    signature_path=signature_path,
                )
            except UniversalFormFillerError:
                pass

        if profile and profile.pdf_coordinates:
            return self._fill_coordinate_pdf(
                reader=reader,
                output_path=output_path,
                data=data,
                profile=profile,
                signature_path=signature_path,
            )

        raise UniversalFormFillerError(
            "PDF has no editable fields and no anchor_rules or pdf_coordinates were provided."
        )

    def _extract_pdf_fields(self, reader: PdfReader) -> Dict[str, Any]:
        try:
            return reader.get_fields() or {}
        except Exception:
            return {}

    def _fill_editable_pdf(
        self,
        reader: PdfReader,
        output_path: str,
        fields: Dict[str, Any],
        data: Dict[str, Any],
        profile: Optional[FormProfile],
        signature_path: Optional[str],
    ) -> Dict[str, Any]:
        writer = PdfWriter()
        resolved_field_map = self._build_pdf_field_map(fields, data, profile)

        for page in reader.pages:
            writer.add_page(page)

        for page_num, page in enumerate(writer.pages):
            page_updates = {}
            for pdf_field_name, field_value in resolved_field_map.items():
                field_pages = self._find_pdf_field_pages(reader, pdf_field_name)
                if page_num in field_pages:
                    page_updates[pdf_field_name] = str(field_value)

            if page_updates:
                writer.update_page_form_field_values(page, page_updates)

        if "/AcroForm" in reader.trailer["/Root"]:
            writer._root_object.update(
                {NameObject("/AcroForm"): reader.trailer["/Root"]["/AcroForm"]}
            )
            try:
                writer._root_object["/AcroForm"].update(
                    {NameObject("/NeedAppearances"): BooleanObject(True)}
                )
            except Exception:
                pass

        with open(output_path, "wb") as f:
            writer.write(f)

        if signature_path and profile and profile.signature_coordinates:
            self._stamp_signature_on_pdf(
                input_pdf=output_path,
                output_pdf=output_path,
                signature_path=signature_path,
                signature_coordinates=profile.signature_coordinates,
            )

        return {
            "type": "editable_pdf",
            "field_count": len(fields),
            "mapped_fields": resolved_field_map,
        }

    def _find_pdf_field_pages(self, reader: PdfReader, target_name: str) -> List[int]:
        matches = []
        for page_index, page in enumerate(reader.pages):
            annotations = page.get("/Annots", [])
            for annot_ref in annotations:
                annot = annot_ref.get_object()
                if annot.get("/T") == target_name:
                    matches.append(page_index)
        return matches

    def _build_pdf_field_map(
        self,
        pdf_fields: Dict[str, Any],
        data: Dict[str, Any],
        profile: Optional[FormProfile],
    ) -> Dict[str, Any]:
        result = {}
        aliases = profile.field_aliases if profile else {}
        normalized_data = {self._normalize_key(k): v for k, v in data.items()}

        for pdf_field_name in pdf_fields.keys():
            normalized_pdf_field = self._normalize_key(pdf_field_name)
            matched_value = None

            if normalized_pdf_field in normalized_data:
                matched_value = normalized_data[normalized_pdf_field]
            else:
                for data_key, alias_list in aliases.items():
                    if normalized_pdf_field in [self._normalize_key(a) for a in alias_list]:
                        data_key_norm = self._normalize_key(data_key)
                        if data_key_norm in normalized_data:
                            matched_value = normalized_data[data_key_norm]
                            break

            if matched_value is None:
                best_key = self._best_key_match(normalized_pdf_field, list(normalized_data.keys()))
                if best_key:
                    matched_value = normalized_data[best_key]

            if matched_value is not None:
                result[pdf_field_name] = matched_value

        return result

    def _fill_anchor_pdf(
        self,
        reader: PdfReader,
        output_path: str,
        data: Dict[str, Any],
        profile: FormProfile,
        signature_path: Optional[str],
    ) -> Dict[str, Any]:
        text_map = self._extract_text_map(reader)
        placements = self._resolve_anchor_placements(text_map=text_map, data=data, profile=profile)

        if not placements and not (signature_path and profile.signature_coordinates):
            raise UniversalFormFillerError("No anchor placements could be resolved from the PDF.")

        self._merge_overlay(
            base_reader=reader,
            output_path=output_path,
            placements=placements,
            signature_path=signature_path,
            signature_coordinates=profile.signature_coordinates,
        )

        return {
            "type": "anchor_pdf",
            "resolved_anchors": [p["field_name"] for p in placements],
        }

    def _fill_coordinate_pdf(
        self,
        reader: PdfReader,
        output_path: str,
        data: Dict[str, Any],
        profile: FormProfile,
        signature_path: Optional[str],
    ) -> Dict[str, Any]:
        placements = []

        for field_name, coords in profile.pdf_coordinates.items():
            if coords.page < 0 or coords.page >= len(reader.pages):
                raise UniversalFormFillerError(
                    f"Profile page index out of range for field '{field_name}': {coords.page}. "
                    f"PDF has {len(reader.pages)} pages indexed 0 to {len(reader.pages) - 1}."
                )

            value = self._resolve_data_value(field_name, data, profile)
            if value is None:
                continue

            placements.append(
                {
                    "field_name": field_name,
                    "page": coords.page,
                    "x": coords.x,
                    "y": coords.y,
                    "width": coords.width,
                    "height": coords.height,
                    "font_size": coords.font_size,
                    "multiline": coords.multiline,
                    "whiteout": coords.whiteout,
                    "whiteout_padding": coords.whiteout_padding,
                    "value": str(value),
                    "max_lines": coords.max_lines,
                    "align": coords.align,
                }
            )

        self._merge_overlay(
            base_reader=reader,
            output_path=output_path,
            placements=placements,
            signature_path=signature_path,
            signature_coordinates=profile.signature_coordinates,
        )

        return {
            "type": "coordinate_pdf",
            "mapped_coordinates": [p["field_name"] for p in placements],
        }

    def _merge_overlay(
        self,
        base_reader: PdfReader,
        output_path: str,
        placements: List[Dict[str, Any]],
        signature_path: Optional[str],
        signature_coordinates: Optional[SignatureCoordinate],
    ) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            overlay_path = os.path.join(tmpdir, "overlay.pdf")
            self._build_pdf_overlay_from_placements(
                base_reader=base_reader,
                overlay_output=overlay_path,
                placements=placements,
                signature_path=signature_path,
                signature_coordinates=signature_coordinates,
            )

            overlay_reader = PdfReader(overlay_path)
            writer = PdfWriter()

            for page_index, base_page in enumerate(base_reader.pages):
                new_page = base_page
                if page_index < len(overlay_reader.pages):
                    new_page.merge_page(overlay_reader.pages[page_index])
                writer.add_page(new_page)

            with open(output_path, "wb") as f:
                writer.write(f)

    def _extract_text_map(self, reader: PdfReader) -> List[Dict[str, Any]]:
        text_map: List[Dict[str, Any]] = []

        for page_index, page in enumerate(reader.pages):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""

            lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
            page_height = float(page.mediabox.height)
            page_width = float(page.mediabox.width)

            total_lines = max(len(lines), 1)
            start_y = page_height - 72
            end_y = 72
            usable_height = max(start_y - end_y, 1)
            line_step = usable_height / total_lines

            for i, line in enumerate(lines):
                y = start_y - (i * line_step)
                text_map.append(
                    {
                        "page": page_index,
                        "text": line,
                        "x": 72.0,
                        "y": y,
                        "page_width": page_width,
                        "page_height": page_height,
                    }
                )

        return text_map

    def _resolve_anchor_placements(
        self,
        text_map: List[Dict[str, Any]],
        data: Dict[str, Any],
        profile: FormProfile,
    ) -> List[Dict[str, Any]]:
        placements: List[Dict[str, Any]] = []

        for rule in profile.anchor_rules:
            value = self._resolve_data_value(rule.field_name, data, profile)
            if value is None:
                continue

            match = self._find_anchor_match(text_map=text_map, rule=rule)
            if not match:
                continue

            placements.append(
                {
                    "field_name": rule.field_name,
                    "page": match["page"],
                    "x": match["x"] + rule.x_offset,
                    "y": match["y"] + rule.y_offset,
                    "width": rule.width,
                    "height": rule.height,
                    "font_size": rule.font_size,
                    "multiline": rule.multiline,
                    "whiteout": rule.whiteout,
                    "whiteout_padding": rule.whiteout_padding,
                    "value": str(value),
                    "max_lines": rule.max_lines,
                    "align": rule.align,
                }
            )

        return placements

    def _find_anchor_match(
        self,
        text_map: List[Dict[str, Any]],
        rule: AnchorRule,
    ) -> Optional[Dict[str, Any]]:
        anchor_norm = self._normalize_text(rule.anchor_text)
        candidates = []

        for item in text_map:
            if rule.page is not None and item["page"] != rule.page:
                continue

            item_text_norm = self._normalize_text(item["text"])

            matched = False
            if rule.match_mode == "exact":
                matched = item_text_norm == anchor_norm
            elif rule.match_mode == "startswith":
                matched = item_text_norm.startswith(anchor_norm)
            else:
                matched = anchor_norm in item_text_norm

            if matched:
                candidates.append(item)

        if not candidates:
            return None

        candidates.sort(key=lambda c: (c["page"], -c["y"]))
        return candidates[0]

    def _build_pdf_overlay_from_placements(
        self,
        base_reader: PdfReader,
        overlay_output: str,
        placements: List[Dict[str, Any]],
        signature_path: Optional[str],
        signature_coordinates: Optional[SignatureCoordinate],
    ) -> None:
        from reportlab.lib.utils import ImageReader

        page_sizes: List[Tuple[float, float]] = []
        for page in base_reader.pages:
            mediabox = page.mediabox
            page_sizes.append((float(mediabox.width), float(mediabox.height)))

        packets = []
        for page_index, (width, height) in enumerate(page_sizes):
            packet_path = f"{overlay_output}.{page_index}.pdf"
            c = canvas.Canvas(packet_path, pagesize=(width, height))

            for placement in placements:
                if placement["page"] != page_index:
                    continue
                self._draw_overlay_value(c, placement)

            if signature_path and signature_coordinates and os.path.exists(signature_path):
                if signature_coordinates.page == page_index:
                    c.drawImage(
                        ImageReader(signature_path),
                        signature_coordinates.x,
                        signature_coordinates.y,
                        width=signature_coordinates.width,
                        height=signature_coordinates.height,
                        preserveAspectRatio=True,
                        mask="auto",
                    )

            c.showPage()
            c.save()
            packets.append(packet_path)

        merged = PdfWriter()
        for page_index, packet in enumerate(packets):
            r = PdfReader(packet)
            if len(r.pages) > 0:
                merged.add_page(r.pages[0])
            else:
                width, height = page_sizes[page_index]
                merged.add_blank_page(width=width, height=height)

        with open(overlay_output, "wb") as f:
            merged.write(f)

    def _draw_overlay_value(self, c: canvas.Canvas, placement: Dict[str, Any]) -> None:
        x = float(placement["x"])
        y = float(placement["y"])
        width = float(placement["width"]) if placement.get("width") is not None else 220.0
        height = float(placement["height"]) if placement.get("height") is not None else 14.0
        font_size = int(placement.get("font_size", 10))
        multiline = bool(placement.get("multiline", False))
        whiteout_enabled = bool(placement.get("whiteout", True))
        whiteout_padding = float(placement.get("whiteout_padding", 2.0))
        value = str(placement.get("value", ""))
        max_lines = placement.get("max_lines")
        align = str(placement.get("align", "left")).lower()

        if whiteout_enabled:
            c.setFillColor(white)
            c.setStrokeColor(white)
            c.rect(
                x - whiteout_padding,
                y - whiteout_padding,
                width + (whiteout_padding * 2),
                height + (whiteout_padding * 2),
                stroke=1,
                fill=1,
            )

        c.setFillColor(black)
        c.setStrokeColor(black)
        c.setFont("Helvetica", font_size)

        if multiline:
            self._draw_multiline_text(
                c=c,
                text=value,
                x=x,
                y=y,
                max_width=width,
                line_height=font_size + 2,
                max_lines=max_lines,
                align=align,
            )
        else:
            if align == "center":
                c.drawCentredString(x + (width / 2.0), y, value)
            elif align == "right":
                c.drawRightString(x + width, y, value)
            else:
                c.drawString(x, y, value)

    def _stamp_signature_on_pdf(
        self,
        input_pdf: str,
        output_pdf: str,
        signature_path: str,
        signature_coordinates: SignatureCoordinate,
    ) -> None:
        reader = PdfReader(input_pdf)

        with tempfile.TemporaryDirectory() as tmpdir:
            overlay_path = os.path.join(tmpdir, "signature_overlay.pdf")
            page_sizes = []
            for page in reader.pages:
                mediabox = page.mediabox
                page_sizes.append((float(mediabox.width), float(mediabox.height)))

            from reportlab.lib.utils import ImageReader

            packets = []
            for page_index, (width, height) in enumerate(page_sizes):
                packet_path = f"{overlay_path}.{page_index}.pdf"
                c = canvas.Canvas(packet_path, pagesize=(width, height))
                if page_index == signature_coordinates.page and os.path.exists(signature_path):
                    c.drawImage(
                        ImageReader(signature_path),
                        signature_coordinates.x,
                        signature_coordinates.y,
                        width=signature_coordinates.width,
                        height=signature_coordinates.height,
                        preserveAspectRatio=True,
                        mask="auto",
                    )
                c.showPage()
                c.save()
                packets.append(packet_path)

            overlay_writer = PdfWriter()
            for page_index, packet in enumerate(packets):
                r = PdfReader(packet)
                if len(r.pages) > 0:
                    overlay_writer.add_page(r.pages[0])
                else:
                    width, height = page_sizes[page_index]
                    overlay_writer.add_blank_page(width=width, height=height)

            with open(overlay_path, "wb") as f:
                overlay_writer.write(f)

            overlay_reader = PdfReader(overlay_path)
            writer = PdfWriter()

            for idx, page in enumerate(reader.pages):
                page.merge_page(overlay_reader.pages[idx])
                writer.add_page(page)

            with open(output_pdf, "wb") as f:
                writer.write(f)

    # ------------------------------------------------------------------
    # DOCX filling
    # ------------------------------------------------------------------

    def _fill_docx(
        self,
        input_path: str,
        output_path: str,
        data: Dict[str, Any],
        profile: Optional[FormProfile],
        signature_path: Optional[str],
    ) -> None:
        doc = Document(input_path)

        self._replace_docx_placeholders(doc, data)

        if signature_path and os.path.exists(signature_path):
            signature_placeholder = (
                profile.docx_signature_placeholder
                if profile and profile.docx_signature_placeholder
                else "{{signature}}"
            )
            self._insert_signature_in_docx(doc, signature_placeholder, signature_path)

        doc.save(output_path)

    def _replace_docx_placeholders(self, doc: Document, data: Dict[str, Any]) -> None:
        for paragraph in doc.paragraphs:
            self._replace_text_in_paragraph(paragraph, data)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self._replace_text_in_paragraph(paragraph, data)

    def _replace_text_in_paragraph(self, paragraph, data: Dict[str, Any]) -> None:
        full_text = "".join(run.text for run in paragraph.runs)
        if not full_text:
            return

        replaced_text = self._replace_placeholders_in_text(full_text, data)
        if replaced_text != full_text:
            for run in paragraph.runs:
                run.text = ""
            if paragraph.runs:
                paragraph.runs[0].text = replaced_text
            else:
                paragraph.add_run(replaced_text)

    def _insert_signature_in_docx(
        self,
        doc: Document,
        signature_placeholder: str,
        signature_path: str,
    ) -> None:
        for paragraph in doc.paragraphs:
            if signature_placeholder in paragraph.text:
                for run in paragraph.runs:
                    if signature_placeholder in run.text:
                        run.text = run.text.replace(signature_placeholder, "")
                paragraph.add_run().add_picture(signature_path, width=Inches(1.6))

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if signature_placeholder in paragraph.text:
                            for run in paragraph.runs:
                                if signature_placeholder in run.text:
                                    run.text = run.text.replace(signature_placeholder, "")
                            paragraph.add_run().add_picture(signature_path, width=Inches(1.6))

    # ------------------------------------------------------------------
    # XLSX filling
    # ------------------------------------------------------------------

    def _fill_xlsx(
        self,
        input_path: str,
        output_path: str,
        data: Dict[str, Any],
        profile: Optional[FormProfile],
        signature_path: Optional[str],
    ) -> None:
        wb = load_workbook(input_path)

        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        new_val = self._replace_placeholders_in_text(cell.value, data)
                        if new_val != cell.value:
                            cell.value = new_val

        if signature_path and os.path.exists(signature_path):
            anchor = profile.xlsx_signature_anchor if profile and profile.xlsx_signature_anchor else None
            if anchor:
                first_sheet = wb.worksheets[0]
                img = XLImage(signature_path)
                img.width = 160
                img.height = 60
                first_sheet.add_image(img, anchor)

        wb.save(output_path)

    # ------------------------------------------------------------------
    # Placeholder helpers
    # ------------------------------------------------------------------

    def _replace_placeholders_in_text(self, text: str, data: Dict[str, Any]) -> str:
        result = text
        for key, value in data.items():
            safe_val = "" if value is None else str(value)
            result = result.replace(f"{{{{{key}}}}}", safe_val)
            result = result.replace(f"[[{key}]]", safe_val)
            result = result.replace(f"<<{key}>>", safe_val)
        return result

    def _resolve_data_value(
        self,
        field_name: str,
        data: Dict[str, Any],
        profile: Optional[FormProfile],
    ) -> Optional[Any]:
        normalized_data = {self._normalize_key(k): v for k, v in data.items()}
        n_field = self._normalize_key(field_name)

        if n_field in normalized_data:
            return normalized_data[n_field]

        if profile:
            aliases = profile.field_aliases.get(field_name, [])
            for alias in aliases:
                n_alias = self._normalize_key(alias)
                if n_alias in normalized_data:
                    return normalized_data[n_alias]

        best = self._best_key_match(n_field, list(normalized_data.keys()))
        return normalized_data.get(best) if best else None

    # ------------------------------------------------------------------
    # Conversion
    # ------------------------------------------------------------------

    def _convert_office_to_pdf(self, input_path: str) -> Optional[str]:
        input_file = Path(input_path)
        if not input_file.exists():
            raise UniversalFormFillerError(f"Cannot convert missing file: {input_path}")

        output_dir = input_file.parent
        cmd = [
            self.libreoffice_binary,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(input_file),
        ]

        try:
            subprocess.run(
                cmd,
                capture_output=True,
                text=True,
                check=True,
            )
        except FileNotFoundError as exc:
            raise UniversalFormFillerError(
                "LibreOffice binary not found. Ensure 'soffice' is installed and on PATH."
            ) from exc
        except subprocess.CalledProcessError as exc:
            raise UniversalFormFillerError(
                f"LibreOffice conversion failed: {exc.stderr or exc.stdout}"
            ) from exc

        pdf_path = input_file.with_suffix(".pdf")
        return str(pdf_path) if pdf_path.exists() else None

    # ------------------------------------------------------------------
    # Utilities
    # ------------------------------------------------------------------

    def _normalize_key(self, value: str) -> str:
        value = value.strip().lower()
        value = re.sub(r"[\s_\-/\.]+", "", value)
        value = re.sub(r"[^a-z0-9]", "", value)
        return value

    def _normalize_text(self, text: str) -> str:
        text = (text or "").lower()
        text = text.replace("&", " and ")
        text = text.replace("’", "'")
        text = text.replace("‘", "'")
        text = re.sub(r"[\r\n\t]+", " ", text)
        text = re.sub(r"[^a-z0-9\.\' ]+", " ", text)
        text = re.sub(r"\s+", " ", text)
        return text.strip()

    def _best_key_match(self, target: str, candidates: List[str]) -> Optional[str]:
        if not target or not candidates:
            return None

        if target in candidates:
            return target

        partials = [c for c in candidates if target in c or c in target]
        if partials:
            return max(partials, key=len)

        best_candidate = None
        best_score = 0
        for c in candidates:
            score = 0
            for ch in set(target):
                if ch in c:
                    score += 1
            if score > best_score:
                best_score = score
                best_candidate = c

        return best_candidate if best_score > 0 else None

    def _draw_multiline_text(
        self,
        c: canvas.Canvas,
        text: str,
        x: float,
        y: float,
        max_width: float,
        line_height: float,
        max_lines: Optional[int] = None,
        align: str = "left",
    ) -> None:
        words = str(text).split()
        lines: List[str] = []
        current_line = ""

        for word in words:
            test_line = f"{current_line} {word}".strip()
            if c.stringWidth(test_line, "Helvetica", c._fontsize) <= max_width:
                current_line = test_line
            else:
                if current_line:
                    lines.append(current_line)
                current_line = word

        if current_line:
            lines.append(current_line)

        if max_lines is not None:
            lines = lines[:max_lines]

        current_y = y
        for line in lines:
            if align == "center":
                c.drawCentredString(x + (max_width / 2.0), current_y, line)
            elif align == "right":
                c.drawRightString(x + max_width, current_y, line)
            else:
                c.drawString(x, current_y, line)
            current_y -= line_height

    def _sanitize_filename(self, name: str) -> str:
        name = re.sub(r"[^\w\-\.]+", "_", name.strip())
        return name[:180] if len(name) > 180 else name


# ------------------------------------------------------------------
# LMCP defaults
# ------------------------------------------------------------------

def build_lmcp_default_form_data(
    tender_data: Optional[Dict[str, Any]] = None,
    company_data: Optional[Dict[str, Any]] = None,
    director_data: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    tender_data = tender_data or {}
    company_data = company_data or {}
    director_data = director_data or {}

    merged = {
        "company_name": company_data.get("company_name", "Lechesa Manaba Consulting and Projects (Pty) Ltd"),
        "trading_name": company_data.get("trading_name", "LMCP"),
        "registration_number": company_data.get("registration_number", ""),
        "vat_number": company_data.get("vat_number", ""),
        "tax_number": company_data.get("tax_number", ""),
        "csd_number": company_data.get("csd_number", ""),
        "cidb_grade": company_data.get("cidb_grade", ""),
        "company_email": company_data.get("company_email", ""),
        "company_phone": company_data.get("company_phone", ""),
        "company_address": company_data.get("company_address", ""),
        "postal_address": company_data.get("postal_address", ""),
        "bank_name": company_data.get("bank_name", ""),
        "bank_account_number": company_data.get("bank_account_number", ""),
        "bank_branch_code": company_data.get("bank_branch_code", ""),
        "director_name": director_data.get("director_name", "Lechesa Manaba"),
        "director_capacity": director_data.get("director_capacity", "Managing Director"),
        "director_id_number": director_data.get("director_id_number", ""),
        "director_email": director_data.get("director_email", ""),
        "director_phone": director_data.get("director_phone", ""),
        "signatory_name": director_data.get("signatory_name", "Lechesa Manaba"),
        "signatory_capacity": director_data.get("signatory_capacity", "Managing Director"),
        "date_signed": tender_data.get("date_signed", ""),
        "tender_number": tender_data.get("tender_number", ""),
        "tender_title": tender_data.get("tender_title", ""),
        "rfq_number": tender_data.get("rfq_number", ""),
        "client_name": tender_data.get("client_name", ""),
        "submission_date": tender_data.get("submission_date", ""),
        "quote_amount": tender_data.get("quote_amount", ""),
        "quote_amount_words": tender_data.get("quote_amount_words", ""),
        "contact_person": tender_data.get("contact_person", ""),
        "contact_email": tender_data.get("contact_email", ""),
        "contact_phone": tender_data.get("contact_phone", ""),
        "project_name": tender_data.get("project_name", ""),
        "project_location": tender_data.get("project_location", ""),
        "name_of_bidder": company_data.get("company_name", "Lechesa Manaba Consulting and Projects (Pty) Ltd"),
        "position": director_data.get("director_capacity", "Managing Director"),
        "signature": "",
    }

    merged.update(company_data)
    merged.update(director_data)
    merged.update(tender_data)
    return merged
