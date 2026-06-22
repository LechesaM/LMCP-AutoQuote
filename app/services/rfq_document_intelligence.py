from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional
from urllib.parse import urlparse
import hashlib
import json
import re

try:
    import requests
except Exception:  # pragma: no cover
    requests = None


ENGINE_VERSION = "RFQ_DOCUMENT_INTELLIGENCE_V1"

RUNTIME_DIR = Path("runtime/rfq_document_intelligence")
DOWNLOAD_DIR = RUNTIME_DIR / "downloads"
REPORT_DIR = RUNTIME_DIR / "reports"

DEFAULT_TIMEOUT_SECONDS = 30
MAX_TEXT_PAGES = 40
MAX_TEXT_CHARS = 250_000
MAX_LOCAL_FILE_BYTES = 50 * 1024 * 1024

PRICING_TERMS = [
    "pricing schedule",
    "price schedule",
    "schedule of prices",
    "schedule of rates",
    "quotation schedule",
    "financial offer",
    "form of offer",
    "activity schedule",
    "pricing data",
    "scope and pricing",
    "gcc pricing schedule",
    "sbd 3.1",
    "sbd 3.2",
    "unit price",
    "total price",
    "sub-total",
    "subtotal",
    "billing model",
    "cost breakdown",
    "rate",
    "amount",
    "rates",
    "price",
    "rate schedule",
]

BOQ_TERMS = [
    "bill of quantities",
    "bills of quantities",
    "boq",
    "pricing schedule",
    "price schedule",
    "schedule of rates",
    "schedule of quantities",
    "schedule of prices",
    "activity schedule",
    "financial offer",
    "form of offer",
    "section c2",
    "pricing data",
    "scope and pricing",
    "gcc pricing schedule",
    "quantity",
    "quantities",
    "unit of measure",
    "uom",
]

RETURNABLE_TERMS = [
    "returnable documents",
    "returnable schedules",
    "mandatory returnables",
    "compulsory returnables",
    "annexure",
    "annexures",
    "appendix",
    "appendices",
    "sbd",
    "tax compliance",
    "central supplier database",
    "csd",
    "b-bbee",
    "bbbee",
]

PRICING_COLUMN_TERMS = [
    "unit price",
    "price",
    "rate",
    "amount",
    "total",
    "sub-total",
    "subtotal",
    "cost",
]

QUANTITY_COLUMN_TERMS = [
    "quantity",
    "qty",
    "uom",
    "unit of measure",
    "unit",
    "quantities",
]

DESCRIPTION_COLUMN_TERMS = [
    "description",
    "item",
    "goods",
    "service",
    "services",
    "activity",
    "scope",
    "pricing data",
]

TECHNICAL_TERMS = [
    "technical specification",
    "scope of work",
    "terms of reference",
    "specification",
    "specifications",
    "delivery",
    "deliver",
    "supply",
]


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _safe_str(value: Any) -> str:
    try:
        if value is None:
            return ""
        return str(value).strip()
    except Exception:
        return ""


def _safe_lower(value: Any) -> str:
    return _safe_str(value).lower()


def _slug(value: Any, limit: int = 100) -> str:
    text = _safe_str(value)
    text = re.sub(r"[^A-Za-z0-9]+", "-", text).strip("-")
    return (text[:limit].strip("-") or "rfq-document")


def _unique_preserve_order(values: List[str]) -> List[str]:
    seen = set()
    out = []
    for value in values:
        cleaned = _safe_str(value)
        if not cleaned or cleaned in seen:
            continue
        seen.add(cleaned)
        out.append(cleaned)
    return out


def _is_http_url(value: str) -> bool:
    try:
        parsed = urlparse(value)
        return parsed.scheme in {"http", "https"} and bool(parsed.netloc)
    except Exception:
        return False


def _extract_local_document_paths(item: Dict[str, Any]) -> List[str]:
    if not isinstance(item, dict):
        return []
    values: List[str] = []
    for key in ("local_document_paths", "document_paths", "uploaded_files", "files"):
        raw = item.get(key)
        if isinstance(raw, list):
            values.extend(_safe_str(value) for value in raw)
        elif raw:
            values.append(_safe_str(raw))
    resolved: List[str] = []
    for value in _unique_preserve_order(values):
        path = Path(value).expanduser()
        if path.exists() and path.is_file():
            resolved.append(str(path.resolve()))
    return resolved


def _read_docx_text(path: Path) -> Dict[str, Any]:
    try:
        from docx import Document
    except Exception as exc:
        return {"status": "failed", "text": "", "tables": [], "error": f"python-docx unavailable: {exc}"}

    try:
        doc = Document(str(path))
    except Exception as exc:
        return {"status": "failed", "text": "", "tables": [], "error": str(exc)}

    paragraphs: List[str] = []
    for paragraph in doc.paragraphs:
        text = re.sub(r"\s+", " ", _safe_str(paragraph.text)).strip()
        if text:
            paragraphs.append(text)

    tables: List[List[List[str]]] = []
    for table in doc.tables:
        rows: List[List[str]] = []
        for row in table.rows:
            cells = [re.sub(r"\s+", " ", _safe_str(cell.text)).strip() for cell in row.cells]
            if any(cells):
                rows.append(cells)
        if rows:
            tables.append(rows)

    text_parts: List[str] = list(paragraphs)
    for table in tables:
        for row in table:
            line = " | ".join(cell for cell in row if cell)
            if line:
                text_parts.append(line)

    return {
        "status": "ok",
        "text": "\n".join(text_parts)[:MAX_TEXT_CHARS],
        "tables": tables,
        "paragraphs": paragraphs,
    }


def _read_spreadsheet_text(path: Path) -> Dict[str, Any]:
    tables: List[List[List[str]]] = []

    try:
        from openpyxl import load_workbook
        workbook = load_workbook(filename=str(path), data_only=True, read_only=True)
    except Exception:
        workbook = None

    if workbook is not None:
        try:
            for sheet in workbook.worksheets:
                rows: List[List[str]] = []
                for row in sheet.iter_rows(values_only=True):
                    cleaned = [re.sub(r"\s+", " ", _safe_str(cell)).strip() for cell in (row or [])]
                    if any(cleaned):
                        rows.append(cleaned)
                if rows:
                    tables.append(rows)
        finally:
            try:
                workbook.close()
            except Exception:
                pass

    if not tables:
        try:
            import xlrd  # type: ignore
            book = xlrd.open_workbook(str(path))
        except Exception:
            book = None

        if book is not None:
            try:
                for sheet in book.sheets():
                    rows: List[List[str]] = []
                    for r in range(sheet.nrows):
                        cleaned = [re.sub(r"\s+", " ", _safe_str(sheet.cell_value(r, c))).strip() for c in range(sheet.ncols)]
                        if any(cleaned):
                            rows.append(cleaned)
                    if rows:
                        tables.append(rows)
            except Exception:
                pass

    text_lines: List[str] = []
    for table in tables:
        for row in table:
            line = " | ".join(cell for cell in row if cell)
            if line:
                text_lines.append(line)

    return {
        "status": "ok" if tables else "empty",
        "text": "\n".join(text_lines)[:MAX_TEXT_CHARS],
        "tables": tables,
    }


def _extract_local_text(path: Path, ext: str) -> Dict[str, Any]:
    ext = ext.lower()
    if ext == ".pdf":
        text = _extract_pdf_text(path)
        if len(text.strip()) < 60:
            text = (text + "\n" + _extract_pdf_table_text(path)).strip()
        return {"status": "ok" if text else "empty", "text": text, "tables": []}
    if ext in {".html", ".htm"}:
        text = _extract_html_text(path)
        return {"status": "ok" if text else "empty", "text": text, "tables": []}
    if ext in {".txt", ".csv"}:
        text = _extract_plain_text(path)
        return {"status": "ok" if text else "empty", "text": text, "tables": []}
    if ext == ".docx":
        return _read_docx_text(path)
    if ext in {".xlsx", ".xlsm", ".xltx", ".xltm", ".xls"}:
        return _read_spreadsheet_text(path)
    return {"status": "unsupported", "text": "", "tables": []}


def _table_rows_blob(tables: List[List[List[str]]]) -> str:
    lines: List[str] = []
    for table in tables:
        for row in table[:25]:
            cells = [_safe_lower(cell) for cell in row if _safe_str(cell)]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _table_has_pricing_headers(tables: List[List[List[str]]]) -> bool:
    for table in tables:
        for row in table[:8]:
            blob = " ".join(_safe_lower(cell) for cell in row if _safe_str(cell))
            if not blob:
                continue
            has_description = any(term in blob for term in DESCRIPTION_COLUMN_TERMS)
            has_quantity = any(term in blob for term in QUANTITY_COLUMN_TERMS)
            has_price = any(term in blob for term in PRICING_COLUMN_TERMS)
            if has_description and (has_quantity or has_price):
                return True
    return False


def _detect_document_roles(text: str, tables: List[List[List[str]]], filename: str = "") -> Dict[str, Any]:
    blob = f"{filename} {text}".lower()
    has_tables = bool(tables)
    table_blob = _table_rows_blob(tables)
    pricing_terms = any(term in blob for term in PRICING_TERMS)
    boq_terms = any(term in blob for term in BOQ_TERMS)
    returnable_terms = any(term in blob for term in RETURNABLE_TERMS)
    annexure_terms = any(term in blob for term in ("annexure", "annexures", "appendix", "appendices"))
    technical_terms = any(term in blob for term in TECHNICAL_TERMS)
    sbd_terms = "sbd" in blob or "standard bidding document" in blob or "declaration of interest" in blob
    pricing_columns = _table_has_pricing_headers(tables) or (
        any(term in table_blob for term in DESCRIPTION_COLUMN_TERMS)
        and any(term in table_blob for term in PRICING_COLUMN_TERMS)
    )
    quantity_columns = any(term in table_blob for term in QUANTITY_COLUMN_TERMS)
    excel_like = any(ext in filename.lower() for ext in (".xlsx", ".xls", ".csv"))
    pricing_data_terms = any(
        term in blob for term in ("pricing data", "section c2", "scope and pricing", "financial offer", "form of offer")
    )

    pricing_schedule_confidence = 0.0
    boq_confidence = 0.0
    returnable_confidence = 0.0
    technical_confidence = 0.0
    pricing_reasons: List[str] = []
    boq_reasons: List[str] = []
    returnable_reasons: List[str] = []
    technical_reasons: List[str] = []

    if excel_like:
        pricing_schedule_confidence += 0.35
        boq_confidence += 0.25
        pricing_reasons.append("spreadsheet_or_csv_file")
        boq_reasons.append("spreadsheet_or_csv_file")
    if has_tables:
        pricing_schedule_confidence += 0.10
        boq_confidence += 0.10
        pricing_reasons.append("has_tables")
        boq_reasons.append("has_tables")
    if pricing_terms:
        pricing_schedule_confidence += 0.30
        pricing_reasons.append("pricing_terms")
    if boq_terms:
        boq_confidence += 0.30
        boq_reasons.append("boq_terms")
    if "bill of quantities" in blob or "bills of quantities" in blob:
        boq_confidence += 0.30
        boq_reasons.append("bill_of_quantities_phrase")
    if pricing_columns:
        pricing_schedule_confidence += 0.25
        pricing_reasons.append("pricing_columns")
        if quantity_columns:
            boq_confidence += 0.20
            boq_reasons.append("quantity_and_pricing_columns")
        else:
            boq_confidence += 0.10
            boq_reasons.append("pricing_table_structure")
    if pricing_data_terms:
        pricing_schedule_confidence += 0.20
        boq_confidence += 0.15
        pricing_reasons.append("pricing_data_section")
        boq_reasons.append("pricing_data_section")
    if any(term in blob for term in ("unit price", "total price", "rate", "amount", "sub-total", "subtotal", "cost")):
        pricing_schedule_confidence += 0.15
        boq_confidence += 0.10
        pricing_reasons.append("price_amount_terms")
        boq_reasons.append("price_amount_terms")
    if sbd_terms:
        returnable_confidence += 0.20
        returnable_reasons.append("sbd_terms")
    if returnable_terms:
        returnable_confidence += 0.35
        returnable_reasons.append("returnable_terms")
    if annexure_terms:
        returnable_confidence += 0.20
        returnable_reasons.append("annexure_terms")
    if "sbd 3.1" in blob or "sbd 3.2" in blob:
        pricing_schedule_confidence += 0.30
        returnable_confidence += 0.15
        pricing_reasons.append("sbd_pricing_form")
        returnable_reasons.append("sbd_pricing_form")
    if technical_terms:
        technical_confidence += 0.35
        technical_reasons.append("technical_terms")

    pricing_schedule_confidence = round(min(pricing_schedule_confidence, 1.0), 4)
    boq_confidence = round(min(boq_confidence, 1.0), 4)
    returnable_confidence = round(min(returnable_confidence, 1.0), 4)
    technical_confidence = round(min(technical_confidence, 1.0), 4)

    role_candidates = {
        "pricing_schedule": pricing_schedule_confidence,
        "boq": boq_confidence,
        "commercial_returnable": returnable_confidence,
        "technical_document": technical_confidence,
    }
    role = max(role_candidates.items(), key=lambda pair: pair[1])[0]
    if role_candidates[role] <= 0.0:
        role = "supporting_document"
    elif role == "pricing_schedule" and pricing_schedule_confidence < 0.65:
        role = "pricing_schedule_shell" if sbd_terms else "commercial_returnable"
    elif role == "boq" and boq_confidence < 0.65:
        role = "boq_shell"
    elif sbd_terms and role not in {"pricing_schedule", "boq"}:
        role = "sbd_form"

    return {
        "document_role": role,
        "pricing_schedule_confidence": pricing_schedule_confidence,
        "boq_confidence": boq_confidence,
        "commercial_returnable_confidence": returnable_confidence,
        "technical_document_confidence": technical_confidence,
        "pricing_schedule_detected": pricing_schedule_confidence >= 0.55,
        "boq_detected": boq_confidence >= 0.55,
        "returnables_detected": returnable_confidence >= 0.45,
        "annexure_detected": annexure_terms,
        "pricing_schedule_reason": ";".join(dict.fromkeys(pricing_reasons)) or ("pricing_terms" if pricing_terms else ""),
        "boq_reason": ";".join(dict.fromkeys(boq_reasons)) or ("boq_terms" if boq_terms else ""),
        "returnables_reason": ";".join(dict.fromkeys(returnable_reasons)) or ("returnable_terms" if returnable_terms else ""),
        "annexure_reason": "annexure_terms" if annexure_terms else "",
        "signals": {
            "pricing_terms": pricing_terms,
            "boq_terms": boq_terms,
            "returnable_terms": returnable_terms,
            "annexure_terms": annexure_terms,
            "technical_terms": technical_terms,
            "sbd_terms": sbd_terms,
            "pricing_columns": pricing_columns,
            "quantity_columns": quantity_columns,
            "pricing_data_terms": pricing_data_terms,
        },
    }


def extract_document_links(item: Dict[str, Any]) -> List[str]:
    """
    Extract candidate document/source URLs from an RFQ item.

    This function intentionally does not crawl deeply. It only uses URLs already
    present on the RFQ object so the engine remains safe and predictable.
    """
    if not isinstance(item, dict):
        return []

    keys = (
        "document_url",
        "source_url",
        "detail_url",
        "pdf_url",
        "download_url",
        "attachment_url",
        "tender_document_url",
    )

    links: List[str] = []

    for key in keys:
        value = item.get(key)
        if isinstance(value, list):
            links.extend(_safe_str(v) for v in value)
        else:
            links.append(_safe_str(value))

    for key in ("documents", "attachments", "supporting_documents", "links"):
        value = item.get(key)
        if isinstance(value, list):
            for entry in value:
                if isinstance(entry, dict):
                    for inner_key in ("url", "href", "download_url", "document_url"):
                        links.append(_safe_str(entry.get(inner_key)))
                else:
                    links.append(_safe_str(entry))
        elif isinstance(value, dict):
            for inner_key in ("url", "href", "download_url", "document_url"):
                links.append(_safe_str(value.get(inner_key)))

    return [u for u in _unique_preserve_order(links) if _is_http_url(u)]


def _guess_extension(url: str, content_type: str = "") -> str:
    lowered_url = _safe_lower(url)
    lowered_type = _safe_lower(content_type)

    if ".pdf" in lowered_url or "application/pdf" in lowered_type or "pdf" in lowered_type:
        return ".pdf"
    if ".xlsx" in lowered_url or "spreadsheet" in lowered_type or "excel" in lowered_type:
        return ".xlsx"
    if ".xls" in lowered_url:
        return ".xls"
    if ".docx" in lowered_url or "wordprocessingml" in lowered_type:
        return ".docx"
    if ".doc" in lowered_url:
        return ".doc"
    if ".csv" in lowered_url or "text/csv" in lowered_type:
        return ".csv"
    if ".txt" in lowered_url or "text/plain" in lowered_type:
        return ".txt"
    return ".html"


def download_rfq_document(url: str, title: str = "", timeout_seconds: int = DEFAULT_TIMEOUT_SECONDS) -> Dict[str, Any]:
    """
    Download a single RFQ document or source page.

    Returns a structured result. It never raises for network failures.
    """
    DOWNLOAD_DIR.mkdir(parents=True, exist_ok=True)

    url = _safe_str(url)
    if not _is_http_url(url):
        return {
            "status": "failed",
            "url": url,
            "error": "invalid_url",
        }

    if requests is None:
        return {
            "status": "failed",
            "url": url,
            "error": "requests_not_available",
        }

    headers = {
        "User-Agent": "LMCP-AutoQuote RFQ Document Intelligence/1.0",
        "Accept": "text/html,application/pdf,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    }

    try:
        response = requests.get(url, timeout=timeout_seconds, headers=headers, allow_redirects=True)
        content_type = response.headers.get("content-type", "")
        status_code = int(response.status_code)

        if status_code >= 400:
            return {
                "status": "failed",
                "url": url,
                "status_code": status_code,
                "content_type": content_type,
                "error": "http_error",
            }

        ext = _guess_extension(str(response.url or url), content_type)
        digest = hashlib.sha256(url.encode("utf-8")).hexdigest()[:14]
        filename = f"{_slug(title or url)}__{digest}{ext}"
        path = DOWNLOAD_DIR / filename
        path.write_bytes(response.content)

        return {
            "status": "downloaded",
            "url": url,
            "final_url": str(response.url),
            "path": str(path),
            "filename": filename,
            "extension": ext,
            "content_type": content_type,
            "status_code": status_code,
            "size_bytes": len(response.content),
        }

    except Exception as exc:
        return {
            "status": "failed",
            "url": url,
            "error": str(exc),
        }


def _extract_pdf_text(path: Path) -> str:
    text = ""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(str(path))
        chunks: List[str] = []
        page_limit = min(len(doc), MAX_TEXT_PAGES)
        for index in range(page_limit):
            page = doc[index]
            chunks.append(page.get_text("text") or "")
        text = "\n".join(chunks)[:MAX_TEXT_CHARS]
    except Exception:
        text = ""

    if len(text.strip()) >= 60:
        return text

    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        try:
            from PyPDF2 import PdfReader  # type: ignore
        except Exception:
            return text

    try:
        reader = PdfReader(str(path))
        chunks = []
        for page in reader.pages[:MAX_TEXT_PAGES]:
            chunks.append(page.extract_text() or "")
        fallback_text = "\n".join(chunks)[:MAX_TEXT_CHARS]
        return fallback_text or text
    except Exception:
        return text


def _extract_pdf_table_text(path: Path) -> str:
    try:
        import pdfplumber  # type: ignore
    except Exception:
        return ""
    lines: List[str] = []
    try:
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages[:MAX_TEXT_PAGES]:
                for table in page.extract_tables() or []:
                    for row in table or []:
                        cells = [_safe_str(cell) for cell in (row or []) if _safe_str(cell)]
                        if cells:
                            lines.append(" | ".join(cells))
        return "\n".join(lines)[:MAX_TEXT_CHARS]
    except Exception:
        return ""


def _extract_html_text(path: Path) -> str:
    try:
        raw = path.read_text(errors="ignore")
    except Exception:
        return ""

    raw = re.sub(r"<script\b[^>]*>.*?</script>", " ", raw, flags=re.I | re.S)
    raw = re.sub(r"<style\b[^>]*>.*?</style>", " ", raw, flags=re.I | re.S)
    raw = re.sub(r"<[^>]+>", " ", raw)
    raw = re.sub(r"&nbsp;?", " ", raw, flags=re.I)
    raw = re.sub(r"&amp;?", "&", raw, flags=re.I)
    raw = re.sub(r"&lt;?", "<", raw, flags=re.I)
    raw = re.sub(r"&gt;?", ">", raw, flags=re.I)
    raw = re.sub(r"\s+", " ", raw).strip()
    return raw[:MAX_TEXT_CHARS]


def _extract_plain_text(path: Path) -> str:
    try:
        return path.read_text(errors="ignore")[:MAX_TEXT_CHARS]
    except Exception:
        return ""


def extract_text_from_download(download: Dict[str, Any]) -> Dict[str, Any]:
    """
    Extract text from a downloaded file where possible.
    """
    if not isinstance(download, dict) or download.get("status") != "downloaded":
        return {
            "status": "skipped",
            "reason": "not_downloaded",
            "text": "",
            "text_length": 0,
        }

    path = Path(_safe_str(download.get("path")))
    if not path.exists():
        return {
            "status": "failed",
            "reason": "file_missing",
            "path": str(path),
            "text": "",
            "text_length": 0,
        }

    ext = _safe_lower(path.suffix)
    if ext == ".pdf":
        text = _extract_pdf_text(path)
        if len(text.strip()) < 60:
            text = (text + "\n" + _extract_pdf_table_text(path)).strip()
        method = "pdf_pymupdf"
    elif ext in {".html", ".htm"}:
        text = _extract_html_text(path)
        method = "html_strip"
    elif ext in {".txt", ".csv"}:
        text = _extract_plain_text(path)
        method = "plain_text"
    elif ext == ".docx":
        extracted = _read_docx_text(path)
        text = _safe_str(extracted.get("text"))
        method = "docx_python_docx"
    elif ext in {".xlsx", ".xlsm", ".xltx", ".xltm", ".xls"}:
        extracted = _read_spreadsheet_text(path)
        text = _safe_str(extracted.get("text"))
        method = "spreadsheet_openpyxl_xlrd"
    else:
        text = ""
        method = "unsupported"

    return {
        "status": "ok" if text else "empty",
        "path": str(path),
        "extension": ext,
        "method": method,
        "text": text,
        "text_length": len(text),
    }


def _find_reference_numbers(text: str) -> List[str]:
    patterns = [
        r"\bRFQ\s*(?:NO|NUMBER|REF|REFERENCE)?\s*[:/#-]?\s*([A-Z0-9][A-Z0-9/&_. -]{3,60})",
        r"\bRFP\s*(?:NO|NUMBER|REF|REFERENCE)?\s*[:/#-]?\s*([A-Z0-9][A-Z0-9/&_. -]{3,60})",
        r"\bBID\s*(?:NO|NUMBER|REF|REFERENCE)?\s*[:/#-]?\s*([A-Z0-9][A-Z0-9/&_. -]{3,60})",
        r"\bTENDER\s*(?:NO|NUMBER|REF|REFERENCE)?\s*[:/#-]?\s*([A-Z0-9][A-Z0-9/&_. -]{3,60})",
        r"\bREF(?:ERENCE)?\s*(?:NO|NUMBER)?\s*[:/#-]?\s*([A-Z0-9][A-Z0-9/&_. -]{3,60})",
    ]

    found: List[str] = []
    for pattern in patterns:
        for match in re.finditer(pattern, text, flags=re.I):
            value = re.sub(r"\s+", " ", match.group(1)).strip(" -:;,.")
            if value and len(value) >= 3:
                found.append(value[:80])

    return _unique_preserve_order(found)[:20]


def _find_dates(text: str) -> List[str]:
    patterns = [
        r"\b\d{1,2}[/-]\d{1,2}[/-]20\d{2}\b",
        r"\b20\d{2}[/-]\d{1,2}[/-]\d{1,2}\b",
        r"\b\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\s+20\d{2}\b",
    ]
    found: List[str] = []
    for pattern in patterns:
        found.extend(re.findall(pattern, text, flags=re.I))
    return _unique_preserve_order(found)[:30]


def _find_quantity_mentions(text: str) -> List[Dict[str, str]]:
    pattern = r"\b(\d+(?:[,.]\d+)?)\s*(each|ea|unit|units|box|boxes|pack|packs|litre|litres|liter|liters|ml|kg|g|ton|tons|metre|metres|meter|meters|m|mm|cm|pairs|sets|rolls|reams)\b"
    matches = re.findall(pattern, text, flags=re.I)
    out = []
    for quantity, unit in matches[:50]:
        out.append({
            "quantity": quantity,
            "unit": unit,
        })
    return out




def _normalise_for_match(value: Any) -> str:
    text = _safe_lower(value)
    text = re.sub(r"&amp;?", "&", text)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _best_title_candidates(item: Dict[str, Any]) -> List[str]:
    if not isinstance(item, dict):
        return []

    raw_candidates = [
        item.get("title"),
        item.get("buyer_rfq_number"),
        item.get("rfq_number"),
        item.get("reference_number"),
        item.get("description"),
    ]

    candidates: List[str] = []
    for value in raw_candidates:
        text = _safe_str(value)
        if not text:
            continue
        candidates.append(text)
        # Use a shorter prefix too because website text may truncate long titles.
        words = text.split()
        if len(words) >= 6:
            candidates.append(" ".join(words[:12]))
            candidates.append(" ".join(words[:8]))

    cleaned = []
    for candidate in candidates:
        norm = _normalise_for_match(candidate)
        if len(norm) >= 15:
            cleaned.append(candidate)

    return _unique_preserve_order(cleaned)


def _find_best_match_index(full_text: str, item: Dict[str, Any]) -> int:
    if not full_text:
        return -1

    full_norm = _normalise_for_match(full_text)
    if not full_norm:
        return -1

    # Build a rough mapping from normalised index to original index.
    # This is approximate but good enough for a nearby window.
    original_lower = _safe_lower(full_text)

    for candidate in _best_title_candidates(item):
        norm_candidate = _normalise_for_match(candidate)
        if not norm_candidate:
            continue

        idx_norm = full_norm.find(norm_candidate)
        if idx_norm >= 0:
            # Try exact lowercase first for better original positioning.
            exact = _safe_lower(candidate)
            exact_idx = original_lower.find(exact)
            if exact_idx >= 0:
                return exact_idx

            # Approximate original index by ratio.
            ratio = idx_norm / max(len(full_norm), 1)
            return int(ratio * len(full_text))

        # Fuzzy fallback: require at least first 5 meaningful words in sequence.
        words = [w for w in norm_candidate.split() if len(w) > 2]
        if len(words) >= 5:
            phrase = " ".join(words[:5])
            idx_norm = full_norm.find(phrase)
            if idx_norm >= 0:
                ratio = idx_norm / max(len(full_norm), 1)
                return int(ratio * len(full_text))

    return -1


def isolate_rfq_text_block(full_text: str, item: Dict[str, Any], before_chars: int = 1500, after_chars: int = 4500) -> Dict[str, Any]:
    """
    Isolate the text block most likely belonging to the current RFQ.

    Many procurement pages are listing pages containing multiple tenders. Analysing
    the whole page can cause false detections, especially compulsory briefing
    notices from unrelated tenders. This function narrows analysis to a local
    window around the current title/RFQ number when possible.
    """
    full_text = _safe_str(full_text)
    if not full_text:
        return {
            "status": "empty",
            "used_isolated_block": False,
            "match_index": -1,
            "text": "",
            "text_length": 0,
            "reason": "empty_text",
        }

    idx = _find_best_match_index(full_text, item)
    if idx < 0:
        return {
            "status": "fallback_full_text",
            "used_isolated_block": False,
            "match_index": -1,
            "text": full_text[:MAX_TEXT_CHARS],
            "text_length": min(len(full_text), MAX_TEXT_CHARS),
            "reason": "title_not_found",
        }

    start = max(0, idx - before_chars)
    end = min(len(full_text), idx + after_chars)
    block = full_text[start:end].strip()

    # Try to trim at obvious next tender boundaries after the match.
    local_match_pos = max(0, idx - start)
    after = block[local_match_pos:]
    boundary_patterns = [
        r"\bBID DESCRIPTION\s*:",
        r"\bTENDER ADVERT\s*:",
        r"\bRETENDER\s*:",
        r"\bRFQ\s+NO\s*:",
        r"\bRFP\s+\d+/",
        r"\bBid Number\s*:",
    ]

    possible_cuts = []
    for pattern in boundary_patterns:
        for m in re.finditer(pattern, after, flags=re.I):
            if m.start() > 250:
                possible_cuts.append(m.start())

    if possible_cuts:
        cut = min(possible_cuts)
        after = after[:cut]

    before = block[:local_match_pos]
    # Keep some leading context, but not too much.
    before = before[-700:]
    block = (before + " " + after).strip()

    return {
        "status": "ok",
        "used_isolated_block": True,
        "match_index": idx,
        "text": block[:MAX_TEXT_CHARS],
        "text_length": len(block[:MAX_TEXT_CHARS]),
        "reason": "matched_title_or_reference",
    }



def analyse_document_text(text: str) -> Dict[str, Any]:
    text = _safe_str(text)
    t = text.lower()

    briefing_terms = [
        "compulsory briefing",
        "mandatory briefing",
        "compulsory site meeting",
        "mandatory site meeting",
        "compulsory clarification meeting",
        "mandatory clarification meeting",
        "non-compulsory briefing",
        "non compulsory briefing",
        "information session",
    ]

    compulsory_briefing_terms = [
        "compulsory briefing",
        "mandatory briefing",
        "compulsory site meeting",
        "mandatory site meeting",
        "compulsory clarification meeting",
        "mandatory clarification meeting",
    ]

    pricing_terms = PRICING_TERMS + BOQ_TERMS + ["quotation schedule"]

    sbd_matches = sorted(set(re.findall(r"\bSBD\s*([0-9](?:\.[0-9])?)\b", text, flags=re.I)))
    pricing_schedule_confidence = 0.0
    boq_confidence = 0.0
    commercial_returnables_confidence = 0.0
    technical_document_confidence = 0.0
    pricing_reasons: List[str] = []
    boq_reasons: List[str] = []
    returnables_reasons: List[str] = []

    if any(term in t for term in PRICING_TERMS):
        pricing_schedule_confidence += 0.35
        pricing_reasons.append("pricing_terms")
    if "unit price" in t or "total price" in t or "rate" in t or "amount" in t or "sub-total" in t or "subtotal" in t:
        pricing_schedule_confidence += 0.20
        boq_confidence += 0.10
        pricing_reasons.append("price_amount_terms")
        boq_reasons.append("price_amount_terms")
    if "boq" in t or "bill of quantities" in t:
        boq_confidence += 0.35
        boq_reasons.append("boq_terms")
    if "bill of quantities" in t or "bills of quantities" in t:
        boq_confidence += 0.25
        boq_reasons.append("bill_of_quantities_phrase")
    if any(term in t for term in ("pricing data", "section c2", "scope and pricing", "activity schedule", "schedule of rates")):
        pricing_schedule_confidence += 0.20
        boq_confidence += 0.20
        pricing_reasons.append("pricing_data_section")
        boq_reasons.append("pricing_data_section")
    if any(term in t for term in RETURNABLE_TERMS):
        commercial_returnables_confidence += 0.35
        returnables_reasons.append("returnable_terms")
    if sbd_matches:
        commercial_returnables_confidence += 0.20
        returnables_reasons.append("sbd_terms")
    if any(term in t for term in ("sbd 3.1", "sbd 3.2")):
        pricing_schedule_confidence += 0.30
        pricing_reasons.append("sbd_pricing_form")
    if any(term in t for term in TECHNICAL_TERMS):
        technical_document_confidence += 0.35
    if "specification" in t or "scope of work" in t or "terms of reference" in t:
        technical_document_confidence += 0.20

    pricing_schedule_confidence = round(min(pricing_schedule_confidence, 1.0), 4)
    boq_confidence = round(min(boq_confidence, 1.0), 4)
    commercial_returnables_confidence = round(min(commercial_returnables_confidence, 1.0), 4)
    technical_document_confidence = round(min(technical_document_confidence, 1.0), 4)

    return {
        "text_length": len(text),
        "has_text": bool(text.strip()),
        "reference_numbers": _find_reference_numbers(text),
        "date_mentions": _find_dates(text),
        "briefing_or_site_meeting_mentioned": any(term in t for term in briefing_terms),
        "compulsory_briefing_required": any(term in t for term in compulsory_briefing_terms),
        "has_pricing_schedule": any(term in t for term in pricing_terms),
        "has_boq": "boq" in t or "bill of quantities" in t,
        "has_sbd_forms": bool(sbd_matches),
        "sbd_forms_detected": [f"SBD {x}" for x in sbd_matches],
        "quantity_mentions": _find_quantity_mentions(text),
        "returnables_mentioned": any(term in t for term in RETURNABLE_TERMS),
        "tax_compliance_mentioned": "tax compliance" in t or "pin" in t and "tax" in t,
        "csd_mentioned": "central supplier database" in t or "csd" in t,
        "bbbee_mentioned": "b-bbee" in t or "bbbee" in t or "bbee" in t,
        "delivery_mentioned": "delivery" in t or "deliver" in t,
        "supply_mentioned": "supply" in t,
        "pricing_schedule_confidence": pricing_schedule_confidence,
        "boq_confidence": boq_confidence,
        "commercial_returnables_confidence": commercial_returnables_confidence,
        "technical_document_confidence": technical_document_confidence,
        "boq_detected": boq_confidence >= 0.55 or "bill of quantities" in t or "boq" in t,
        "pricing_schedule_detected": pricing_schedule_confidence >= 0.55,
        "returnables_detected": commercial_returnables_confidence >= 0.45,
        "boq_detection_reason": ";".join(dict.fromkeys(boq_reasons)),
        "pricing_schedule_detection_reason": ";".join(dict.fromkeys(pricing_reasons)),
        "returnables_detection_reason": ";".join(dict.fromkeys(returnables_reasons)),
    }


def analyse_rfq_documents(item: Dict[str, Any], timeout_seconds: int = DEFAULT_TIMEOUT_SECONDS) -> Dict[str, Any]:
    """
    Main entry point for RFQ Document Intelligence.

    Downloads URLs already present in the RFQ item, extracts readable text,
    detects pricing/BOQ/SBD/briefing/quantity signals, and writes a JSON report.
    """
    REPORT_DIR.mkdir(parents=True, exist_ok=True)

    if not isinstance(item, dict):
        item = {}

    title = _safe_str(
        item.get("title")
        or item.get("buyer_rfq_number")
        or item.get("rfq_number")
        or "rfq"
    )

    links = extract_document_links(item)
    downloads = [download_rfq_document(url, title, timeout_seconds=timeout_seconds) for url in links]
    for local_path in _extract_local_document_paths(item):
        path = Path(local_path)
        if not path.exists() or path.stat().st_size > MAX_LOCAL_FILE_BYTES:
            continue
        downloads.append(
            {
                "status": "downloaded",
                "url": f"file://{path}",
                "final_url": f"file://{path}",
                "path": str(path),
                "filename": path.name,
                "extension": path.suffix.lower(),
                "content_type": "",
                "status_code": 200,
                "size_bytes": path.stat().st_size,
                "local_file": True,
            }
        )

    extracted_text_results: List[Dict[str, Any]] = []
    combined_text_parts: List[str] = []
    document_classifications: List[Dict[str, Any]] = []

    for download in downloads:
        text_result = extract_text_from_download(download)
        text = _safe_str(text_result.get("text"))
        clean_text_result = dict(text_result)
        clean_text_result.pop("text", None)
        extracted_text_results.append(clean_text_result)
        classification = _detect_document_roles(
            text,
            [],
            filename=_safe_str(download.get("filename") or Path(_safe_str(download.get("path"))).name),
        )
        document_classifications.append({
            "filename": _safe_str(download.get("filename") or Path(_safe_str(download.get("path"))).name),
            "path": _safe_str(download.get("path")),
            "extension": _safe_str(download.get("extension")),
            "status": _safe_str(download.get("status")),
            **classification,
        })
        if text:
            combined_text_parts.append(text)

    fallback_text = " ".join([
        _safe_str(item.get("title")),
        _safe_str(item.get("description")),
        _safe_str(item.get("raw_text")),
        _safe_str(item.get("buyer_rfq_number")),
    ]).strip()

    combined_text = "\n\n".join(combined_text_parts).strip()
    if not combined_text and fallback_text:
        combined_text = fallback_text

    isolation = isolate_rfq_text_block(combined_text, item)
    analysis_text = _safe_str(isolation.get("text")) or combined_text

    intelligence = analyse_document_text(analysis_text)
    pricing_schedule_confidence = round(max([float(row.get("pricing_schedule_confidence") or 0.0) for row in document_classifications] or [0.0]), 4)
    boq_confidence = round(max([float(row.get("boq_confidence") or 0.0) for row in document_classifications] or [0.0]), 4)
    commercial_returnables_confidence = round(max([float(row.get("commercial_returnable_confidence") or 0.0) for row in document_classifications] or [0.0]), 4)
    technical_document_confidence = round(max([float(row.get("technical_document_confidence") or 0.0) for row in document_classifications] or [0.0]), 4)
    pricing_schedule_files = [row for row in document_classifications if row.get("document_role") in {"pricing_schedule", "pricing_schedule_shell"}]
    boq_files = [row for row in document_classifications if row.get("document_role") in {"boq", "boq_shell"}]
    commercial_returnable_files = [row for row in document_classifications if row.get("document_role") in {"commercial_returnable", "sbd_form"}]
    annexure_files = [
        row
        for row in document_classifications
        if "annexure_files" in (row.get("categories") if isinstance(row.get("categories"), list) else [])
        or "annexure" in _safe_str(row.get("annexure_reason")).lower()
    ]
    technical_files = [row for row in document_classifications if row.get("document_role") == "technical_document"]

    intelligence["pricing_schedule_confidence"] = max(float(intelligence.get("pricing_schedule_confidence") or 0.0), pricing_schedule_confidence)
    intelligence["boq_confidence"] = max(float(intelligence.get("boq_confidence") or 0.0), boq_confidence)
    intelligence["commercial_returnables_confidence"] = max(float(intelligence.get("commercial_returnables_confidence") or 0.0), commercial_returnables_confidence)
    intelligence["technical_document_confidence"] = max(float(intelligence.get("technical_document_confidence") or 0.0), technical_document_confidence)
    intelligence["pricing_schedule_detected"] = bool(intelligence.get("pricing_schedule_detected") or intelligence["pricing_schedule_confidence"] >= 0.55 or pricing_schedule_files)
    intelligence["boq_detected"] = bool(intelligence.get("boq_detected") or intelligence["boq_confidence"] >= 0.55 or boq_files)
    intelligence["returnables_detected"] = bool(
        intelligence.get("returnables_detected") or intelligence["commercial_returnables_confidence"] >= 0.45 or commercial_returnable_files or annexure_files
    )
    intelligence["pricing_schedule_detection_confidence"] = intelligence["pricing_schedule_confidence"]
    intelligence["boq_detection_confidence"] = intelligence["boq_confidence"]
    intelligence["returnables_detection_confidence"] = intelligence["commercial_returnables_confidence"]
    intelligence["annexure_detection_confidence"] = round(
        max([float(row.get("commercial_returnable_confidence") or row.get("score") or 0.0) for row in annexure_files] or [0.0]),
        4,
    )
    intelligence["pricing_schedule_detection_reason"] = next(
        (_safe_str(row.get("pricing_schedule_reason")) for row in pricing_schedule_files if _safe_str(row.get("pricing_schedule_reason"))),
        _safe_str(intelligence.get("pricing_schedule_detection_reason")),
    )
    intelligence["boq_detection_reason"] = next(
        (_safe_str(row.get("boq_reason")) for row in boq_files if _safe_str(row.get("boq_reason"))),
        _safe_str(intelligence.get("boq_detection_reason")),
    )
    intelligence["returnables_detection_reason"] = next(
        (_safe_str(row.get("returnables_reason")) for row in commercial_returnable_files if _safe_str(row.get("returnables_reason"))),
        _safe_str(intelligence.get("returnables_detection_reason")),
    )
    intelligence["annexure_detection_reason"] = next(
        (_safe_str(row.get("annexure_reason")) for row in annexure_files if _safe_str(row.get("annexure_reason"))),
        _safe_str(intelligence.get("annexure_detection_reason")),
    )
    intelligence["document_classifications"] = document_classifications
    intelligence["pricing_schedule_files"] = pricing_schedule_files
    intelligence["boq_files"] = boq_files
    intelligence["commercial_returnable_files"] = commercial_returnable_files
    intelligence["annexure_files"] = annexure_files
    intelligence["technical_files"] = technical_files
    inventory_paths = [_safe_str(row.get("path")) for row in document_classifications if _safe_str(row.get("path"))]
    intelligence["artifact_count"] = len(inventory_paths)
    intelligence["pdf_count"] = sum(1 for row in document_classifications if _safe_str(row.get("extension")) == ".pdf")
    intelligence["docx_count"] = sum(1 for row in document_classifications if _safe_str(row.get("extension")) == ".docx")
    intelligence["xlsx_count"] = sum(1 for row in document_classifications if _safe_str(row.get("extension")) in {".xlsx", ".xls", ".xlsm", ".xltx", ".xltm"})
    intelligence["zip_count"] = sum(1 for row in downloads if _safe_str(row.get("extension")) == ".zip")
    intelligence["csv_count"] = sum(1 for row in document_classifications if _safe_str(row.get("extension")) == ".csv")
    intelligence["extracted_file_count"] = len(document_classifications)
    intelligence["main_document_path"] = next(
        (
            _safe_str(row.get("path"))
            for row in document_classifications
            if _safe_str(row.get("filename")).lower().startswith("1. main")
            or "main document" in _safe_str(row.get("filename")).lower()
        ),
        "",
    )
    intelligence["detected_document_types"] = sorted(
        dict.fromkeys(
            _safe_str(row.get("extension")).lstrip(".")
            for row in document_classifications
            if _safe_str(row.get("extension"))
        )
    )
    if annexure_files:
        intelligence["detected_document_types"] = sorted(
            dict.fromkeys(list(intelligence["detected_document_types"]) + ["annexure"])
        )
    intelligence["document_inventory_paths_limited"] = inventory_paths[:25]
    intelligence["scanned_pdf_likely"] = bool(downloads) and not bool(analysis_text.strip())
    intelligence["extraction_confidence"] = round(
        min(
            0.98,
            (0.2 if downloads else 0.0)
            + (0.35 if intelligence.get("has_text") else 0.0)
            + (0.15 if intelligence.get("has_pricing_schedule") else 0.0)
            + (0.15 if intelligence.get("quantity_mentions") else 0.0)
            + (0.15 if intelligence.get("reference_numbers") else 0.0),
        ),
        4,
    )
    intelligence["rfq_block_isolation"] = {
        "status": isolation.get("status"),
        "used_isolated_block": isolation.get("used_isolated_block"),
        "match_index": isolation.get("match_index"),
        "text_length": isolation.get("text_length"),
        "reason": isolation.get("reason"),
    }

    quote_safe = not intelligence.get("compulsory_briefing_required", False)

    result: Dict[str, Any] = {
        "status": "ok",
        "engine_version": ENGINE_VERSION,
        "analysed_at": _now_iso(),
        "title": title,
        "buyer_name": _safe_str(item.get("buyer_name")),
        "buyer_rfq_number": _safe_str(item.get("buyer_rfq_number")),
        "links": links,
        "downloads": downloads,
        "text_extraction": extracted_text_results,
        "document_classifications": document_classifications,
        "document_intelligence": intelligence,
        "quote_safe": quote_safe,
        "quote_block_reason": "" if quote_safe else "compulsory_briefing_or_site_meeting_detected",
    }

    report_name = f"{_slug(title)}__document_intelligence.json"
    report_path = REPORT_DIR / report_name
    report_path.write_text(json.dumps(result, indent=2, default=str))
    result["report_path"] = str(report_path)

    return result


def analyse_rfq_document_intelligence(item: Dict[str, Any], timeout_seconds: int = DEFAULT_TIMEOUT_SECONDS) -> Dict[str, Any]:
    """
    Alias kept for future API naming compatibility.
    """
    return analyse_rfq_documents(item, timeout_seconds=timeout_seconds)


def get_status() -> Dict[str, Any]:
    return {
        "status": "ok",
        "engine_version": ENGINE_VERSION,
        "runtime_dir": str(RUNTIME_DIR),
        "download_dir": str(DOWNLOAD_DIR),
        "report_dir": str(REPORT_DIR),
        "capabilities": [
            "download_document_source_detail_urls",
            "extract_pdf_text_with_pymupdf_when_available",
            "extract_html_text",
            "detect_compulsory_briefing_or_site_meeting",
            "detect_pricing_schedule_or_boq",
            "detect_sbd_forms",
            "detect_quantity_mentions",
            "write_json_report",
        ],
    }
