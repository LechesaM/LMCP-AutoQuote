from __future__ import annotations
RETURNABLE_LINE_ITEM_REJECT_TERMS = [
    "bidder to submit",
    "submit reference letters",
    "reference letters",
    "purchase orders",
    "client letter head",
    "client's letter head",
    "clients letter head",
    "contact details",
    "within the last",
    "proof of",
    "mandatory returnable",
    "returnable document",
    "returnable documents",
    "company registration",
    "tax compliance",
    "b-bbee",
    "bbbee",
    "bbee",
    "csd",
    "central supplier database",
    "declaration of interest",
    "sbd",
]


from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
import json
import re

from app.services.rfq_requirement_pack_service import (
    attach_requirement_pack_fields,
    build_requirement_pack,
    normalize_requirement_rows,
)


ENGINE_VERSION = "RFQ_DOCX_MAIN_DOCUMENT_INTELLIGENCE_V1"

RUNTIME_DIR = Path("runtime/rfq_docx_main_document_intelligence")
REPORT_DIR = RUNTIME_DIR / "reports"

MAX_TEXT_CHARS = 600_000

KNOWN_UNITS = {
    "each", "ea", "unit", "units", "box", "boxes", "pack", "packs",
    "lot", "lots", "set", "sets", "pair", "pairs", "roll", "rolls",
    "ream", "reams", "kg", "g", "ton", "tons", "litre", "litres",
    "liter", "liters", "ml", "m", "mm", "cm", "metre", "metres",
    "meter", "meters", "drum", "drums", "container", "containers",
}

PRICING_TERMS = [
    "pricing schedule",
    "price schedule",
    "schedule of prices",
    "quotation schedule",
    "financial offer",
    "form of offer",
    "unit price",
    "total price",
    "amount",
]

BOQ_TERMS = [
    "bill of quantities",
    "boq",
    "pricing schedule",
    "schedule of quantities",
    "schedule of prices",
    "quantity",
    "unit of measure",
    "uom",
]

BRIEFING_TERMS = [
    "compulsory briefing",
    "mandatory briefing",
    "compulsory site meeting",
    "mandatory site meeting",
    "compulsory clarification meeting",
    "mandatory clarification meeting",
]

NON_COMPULSORY_BRIEFING_TERMS = [
    "non-compulsory briefing",
    "non compulsory briefing",
    "not compulsory",
    "optional briefing",
]

DELIVERY_TERMS = [
    "delivery",
    "deliver",
    "offloading",
    "off-loading",
    "delivery address",
    "place of delivery",
]

SPEC_SECTION_TERMS = [
    "scope of work",
    "specification",
    "specifications",
    "technical specification",
    "terms of reference",
    "description of goods",
    "bid description",
]


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _safe_str(value: Any) -> str:
    try:
        if value is None:
            return ""
        return str(value).strip()
    except Exception:
        return ""


def _safe_lower(value: Any) -> str:
    return _safe_str(value).lower()


def _slug(value: Any, limit: int = 100) -> str:
    text = _safe_str(value)
    text = re.sub(r"[^A-Za-z0-9]+", "-", text).strip("-")
    return (text[:limit].strip("-") or "docx-main-document")


def _ensure_dirs() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)


def _clean_text(text: Any) -> str:
    text = _safe_str(text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _to_float(value: Any) -> Optional[float]:
    raw = _safe_str(value)
    if not raw:
        return None
    raw = raw.replace(",", "")
    raw = re.sub(r"[^0-9.\-]", "", raw)
    if raw in {"", ".", "-", "-."}:
        return None
    try:
        return float(raw)
    except Exception:
        return None


def _read_docx(path: Path) -> Dict[str, Any]:
    try:
        from docx import Document
    except Exception as exc:
        return {
            "status": "failed",
            "error": f"python-docx unavailable: {exc}",
            "paragraphs": [],
            "tables": [],
            "text": "",
        }

    try:
        doc = Document(str(path))
    except Exception as exc:
        return {
            "status": "failed",
            "error": str(exc),
            "paragraphs": [],
            "tables": [],
            "text": "",
        }

    paragraphs: List[str] = []
    for paragraph in doc.paragraphs:
        text = _clean_text(paragraph.text)
        if text:
            paragraphs.append(text)

    tables: List[List[List[str]]] = []
    for table in doc.tables:
        rows: List[List[str]] = []
        for row in table.rows:
            cells = [_clean_text(cell.text) for cell in row.cells]
            if any(cells):
                rows.append(cells)
        if rows:
            tables.append(rows)

    table_text_parts: List[str] = []
    for table in tables:
        for row in table:
            table_text_parts.append(" | ".join(row))

    full_text = "\n".join(paragraphs + table_text_parts)
    full_text = full_text[:MAX_TEXT_CHARS]

    return {
        "status": "ok",
        "error": "",
        "paragraphs": paragraphs,
        "tables": tables,
        "text": full_text,
    }


def _find_bid_reference(text: str) -> str:
    patterns = [
        r"\bFIN[-\s_/]*SCM[-\s_/]*(?:TEN|RFQ|BID|EOI)[-\s_/]*\d{3,6}\b",
        r"\bSCM/[A-Z0-9]{2,10}/\d{2,4}\b",
        r"\b[A-Z]-[A-Z]{2}\s*\d{1,4}\s*[-/]\s*20\d{2}\b",
        r"\bT\d{1,3}/\d{1,3}/\d{2,4}\b",
        r"\bRFQ[-\s_/]*[A-Z0-9][A-Z0-9\-_/]{3,30}\b",
        r"\bRFP[-\s_/]*[A-Z0-9][A-Z0-9\-_/]{3,30}\b",
        r"\bBID\s*(?:NO|NUMBER)?\s*[:#-]?\s*([A-Z0-9][A-Z0-9\-_/ ]{3,50})",
        r"\bTENDER\s*(?:NO|NUMBER)?\s*[:#-]?\s*([A-Z0-9][A-Z0-9\-_/ ]{3,50})",
    ]

    for pattern in patterns:
        match = re.search(pattern, text, flags=re.I)
        if match:
            value = match.group(1) if match.groups() else match.group(0)
            return _clean_text(value).strip(" .,:;")
    return ""


def _find_closing_date(text: str) -> str:
    closing_patterns = [
        r"(?:closing date|bid closing date|tender closing date|closing)\s*[:\-]?\s*([0-3]?\d\s+(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\s+20\d{2}(?:\s+at\s+\d{1,2}:\d{2})?)",
        r"(?:closing date|bid closing date|tender closing date|closing)\s*[:\-]?\s*(\d{1,2}[/-]\d{1,2}[/-]20\d{2}(?:\s+\d{1,2}:\d{2})?)",
        r"(?:closing date|bid closing date|tender closing date|closing)\s*[:\-]?\s*(20\d{2}[/-]\d{1,2}[/-]\d{1,2}(?:\s+\d{1,2}:\d{2})?)",
    ]

    for pattern in closing_patterns:
        match = re.search(pattern, text, flags=re.I)
        if match:
            return _clean_text(match.group(1))

    # Fallback to any date only if close to the word closing.
    lower = text.lower()
    idx = lower.find("closing")
    if idx >= 0:
        window = text[idx:idx + 300]
        generic = re.search(
            r"\b(\d{1,2}[/-]\d{1,2}[/-]20\d{2}|20\d{2}[/-]\d{1,2}[/-]\d{1,2}|\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\s+20\d{2})\b",
            window,
            flags=re.I,
        )
        if generic:
            return _clean_text(generic.group(1))

    return ""


def _find_bid_description(text: str, paragraphs: List[str]) -> str:
    patterns = [
        r"(?:bid description|description of bid|tender description|rfq description)\s*[:\-]?\s*(.{20,400})",
        r"(?:for the supply and delivery|supply and delivery|supply, delivery)(.{10,300})",
    ]

    for pattern in patterns:
        match = re.search(pattern, text, flags=re.I | re.S)
        if match:
            value = _clean_text(match.group(1))
            value = re.split(r"\b(?:closing date|validity period|bid number|tender number)\b", value, flags=re.I)[0]
            if len(value) >= 15:
                if pattern.startswith("(?:for"):
                    return _clean_text("Supply and delivery" + value)
                return value[:500]

    for paragraph in paragraphs:
        p = _safe_lower(paragraph)
        if any(term in p for term in ["supply and delivery", "supply, delivery", "bid description"]):
            return paragraph[:500]

    return ""


def _has_pricing_schedule(text: str, tables: List[List[List[str]]]) -> bool:
    t = _safe_lower(text)
    if any(term in t for term in PRICING_TERMS):
        return True
    for table in tables:
        if _table_is_pricing_or_boq(table):
            return True
    return False


def _has_boq_table(text: str, tables: List[List[List[str]]]) -> bool:
    t = _safe_lower(text)
    if "bill of quantities" in t or "boq" in t:
        return True
    return any(_table_is_pricing_or_boq(table) for table in tables)


def _briefing_required(text: str) -> bool:
    t = _safe_lower(text)
    if any(term in t for term in NON_COMPULSORY_BRIEFING_TERMS):
        # If document explicitly says non-compulsory, do not block unless there is
        # another explicit compulsory phrase elsewhere.
        t_without_non_compulsory = t
        for term in NON_COMPULSORY_BRIEFING_TERMS:
            t_without_non_compulsory = t_without_non_compulsory.replace(term, "")
        return any(term in t_without_non_compulsory for term in BRIEFING_TERMS)
    return any(term in t for term in BRIEFING_TERMS)


def _delivery_requirements(text: str, paragraphs: List[str]) -> List[str]:
    found: List[str] = []
    for paragraph in paragraphs:
        p = _safe_lower(paragraph)
        if any(term in p for term in DELIVERY_TERMS):
            found.append(paragraph)
        if len(found) >= 20:
            break
    if found:
        return found

    # Fallback windows
    lower = text.lower()
    for term in DELIVERY_TERMS:
        idx = lower.find(term)
        if idx >= 0:
            found.append(_clean_text(text[max(0, idx - 100):idx + 300]))
            break
    return found


def _scope_or_spec_sections(paragraphs: List[str]) -> List[str]:
    sections: List[str] = []
    capture = False
    current: List[str] = []

    heading_terms = SPEC_SECTION_TERMS
    stop_terms = [
        "pricing schedule",
        "evaluation criteria",
        "returnable documents",
        "closing date",
        "sbd",
        "declaration",
    ]

    for paragraph in paragraphs:
        p = _safe_lower(paragraph)
        if any(term in p for term in heading_terms):
            if current:
                sections.append("\n".join(current))
                current = []
            capture = True
            current.append(paragraph)
            continue

        if capture and any(term in p for term in stop_terms):
            if current:
                sections.append("\n".join(current))
                current = []
            capture = False
            continue

        if capture:
            current.append(paragraph)
            if len(current) >= 25:
                sections.append("\n".join(current))
                current = []
                capture = False

    if current:
        sections.append("\n".join(current))

    return sections[:10]


def _table_is_pricing_or_boq(table: List[List[str]]) -> bool:
    for row in table[:8]:
        joined = " ".join(_safe_lower(cell) for cell in row)
        has_desc = any(x in joined for x in ["description", "item", "specification", "goods", "service"])
        has_qty = any(x in joined for x in ["quantity", "qty", "unit", "uom", "unit of measure"])
        has_price = any(x in joined for x in ["unit price", "total price", "amount", "rate", "price"])
        if has_desc and (has_qty or has_price):
            return True
    return False


def _column_indexes(header: List[str]) -> Dict[str, int]:
    indexes: Dict[str, int] = {}
    for idx, cell in enumerate(header):
        c = _safe_lower(cell)

        if "material number" in c or "material no" in c or c == "material":
            indexes.setdefault("material_number", idx)

        if "item code" in c or c in {"code", "item no", "item number", "no"}:
            indexes.setdefault("item_code", idx)

        if "part number" in c or "part no" in c:
            indexes.setdefault("part_number", idx)

        if "description" in c or "specification" in c or "goods" in c or "service" in c or "item" == c:
            indexes.setdefault("description", idx)

        if "qty" in c or "quantity" in c:
            indexes.setdefault("quantity", idx)

        if c in {"unit", "uom"} or "unit of measure" in c:
            indexes.setdefault("unit", idx)

        if "specification" in c:
            indexes.setdefault("specification", idx)

        if "unit price" in c or "unit rate" in c or c == "rate" or "price quoted" in c:
            indexes.setdefault("buyer_rate", idx)

        if "total price" in c or "total amount" in c or c == "amount":
            indexes.setdefault("buyer_amount", idx)

    return indexes


def _extract_table_line_items(tables: List[List[List[str]]]) -> List[Dict[str, Any]]:
    items: List[Dict[str, Any]] = []

    for table_index, table in enumerate(tables):
        if not table or not _table_is_pricing_or_boq(table):
            continue

        header_index = -1
        for i, row in enumerate(table[:8]):
            if _table_is_pricing_or_boq([row]):
                header_index = i
                break

        if header_index < 0:
            continue

        header = table[header_index]
        indexes = _column_indexes(header)

        if "description" not in indexes:
            # Try fallback if first two columns look like item + description.
            if len(header) >= 2:
                indexes.setdefault("item_code", 0)
                indexes.setdefault("description", 1)
            else:
                continue

        for row_index, row in enumerate(table[header_index + 1:], start=header_index + 1):
            item = _row_to_item(row, indexes, source=f"docx_table_{table_index}")
            if item:
                item["source_table"] = f"docx_table_{table_index}"
                item["source_row"] = row_index
                items.append(item)

    return _dedupe_items(items)




def _is_returnable_or_compliance_line_item(description: str) -> bool:
    desc_l = _safe_lower(description)
    if not desc_l:
        return False
    return any(term in desc_l for term in RETURNABLE_LINE_ITEM_REJECT_TERMS)

def _row_to_item(row: List[str], indexes: Dict[str, int], source: str) -> Optional[Dict[str, Any]]:
    def get(name: str) -> str:
        idx = indexes.get(name)
        if idx is None or idx >= len(row):
            return ""
        return _clean_text(row[idx])

    description = get("description")
    specification = get("specification")
    item_code = get("item_code")
    material_number = get("material_number")
    part_number = get("part_number")
    quantity_raw = get("quantity")
    unit = get("unit")
    buyer_rate = _to_float(get("buyer_rate"))
    buyer_amount = _to_float(get("buyer_amount"))

    if not description and specification:
        description = specification

    if not description:
        joined = " ".join(_clean_text(c) for c in row)
        if len(joined) >= 15:
            description = joined[:300]

    quantity = _to_float(quantity_raw)

    if quantity is None:
        # Conservative row fallback: only accept if row has a clear quantity cell
        # next to a known unit. This avoids hallucinating from dates/durations.
        for i, cell in enumerate(row):
            q = _to_float(cell)
            if q is None:
                continue
            nearby = " ".join(row[max(0, i - 1): min(len(row), i + 2)]).lower()
            if any(re.search(rf"\b{re.escape(unit)}\b", nearby) for unit in KNOWN_UNITS):
                quantity = q
                break

    if not unit:
        joined = " ".join(row)
        unit_pattern = "|".join(re.escape(u) for u in sorted(KNOWN_UNITS, key=len, reverse=True))
        m = re.search(rf"\b(?:\d+(?:[,.]\d+)?)\s*({unit_pattern})\b", joined, flags=re.I)
        if m:
            unit = m.group(1)

    desc_l = _safe_lower(description)
    if any(term in desc_l for term in ["subtotal", "total", "vat", "signature", "name of bidder", "company name"]):
        return None

    if _is_returnable_or_compliance_line_item(description):
        return None

    if quantity is not None and not unit:
        unit = "Each"

    confidence = 0.65
    evidence = ["docx_table"]

    if quantity is not None:
        confidence += 0.15
        evidence.append("quantity_from_table")
    if unit:
        confidence += 0.10
        evidence.append("unit_from_table")
    if item_code:
        confidence += 0.05
        evidence.append("item_code_from_table")
    if specification:
        confidence += 0.05
        evidence.append("specification_from_table")

    return {
        "description": description,
        "quantity": float(quantity) if quantity is not None else None,
        "unit": unit,
        "item_code": item_code,
        "material_number": material_number,
        "part_number": part_number,
        "specification": specification if specification != description else "",
        "buyer_rate": buyer_rate,
        "buyer_amount": buyer_amount,
        "confidence": round(min(confidence, 0.95), 4),
        "evidence": evidence,
        "source": source,
    }


def _extract_explicit_quantity_rows(paragraphs: List[str]) -> List[Dict[str, Any]]:
    """
    Conservative fallback: only extract explicit rows like:
    Quantity: 10 each
    Required quantity 20 drums
    This does not scan arbitrary prose for dates/durations.
    """
    items: List[Dict[str, Any]] = []
    unit_pattern = "|".join(re.escape(u) for u in sorted(KNOWN_UNITS, key=len, reverse=True))

    for paragraph in paragraphs:
        p = _safe_lower(paragraph)
        if not any(signal in p for signal in ["quantity:", "qty:", "required quantity", "number required"]):
            continue

        match = re.search(rf"(?:quantity|qty|required quantity|number required)\s*[:\-]?\s*(\d+(?:[,.]\d+)?)\s*({unit_pattern})\b", paragraph, flags=re.I)
        if not match:
            continue

        quantity = _to_float(match.group(1))
        unit = match.group(2)

        if quantity is None:
            continue

        description = paragraph
        description = re.sub(rf"(?:quantity|qty|required quantity|number required)\s*[:\-]?\s*\d+(?:[,.]\d+)?\s*{re.escape(unit)}", "", description, flags=re.I)
        description = _clean_text(description)

        if len(description) < 8:
            description = "Explicit quantity row from main document"

        items.append({
            "description": description[:300],
            "quantity": float(quantity),
            "unit": unit,
            "item_code": "",
            "specification": "",
            "confidence": 0.70,
            "evidence": ["explicit_quantity_row"],
            "source": "docx_paragraph_explicit_quantity",
        })

    return items[:50]


def _dedupe_items(items: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    seen = set()

    for item in items:
        desc = _safe_lower(item.get("description"))[:160]
        qty = item.get("quantity")
        unit = _safe_lower(item.get("unit"))
        code = _safe_lower(item.get("item_code"))
        key = (desc, qty, unit, code)
        if key in seen:
            continue
        seen.add(key)
        out.append(item)

    return out


def _serialise_tables(tables: List[List[List[str]]]) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    for idx, table in enumerate(tables):
        out.append({
            "index": idx,
            "row_count": len(table),
            "column_count": max((len(row) for row in table), default=0),
            "is_pricing_or_boq": _table_is_pricing_or_boq(table),
            "rows": table,
        })
    return out


def _overall_confidence(
    text: str,
    tables: List[List[List[str]]],
    line_items: List[Dict[str, Any]],
    has_pricing_schedule: bool,
    has_boq_table: bool,
) -> float:
    confidence = 0.0

    if text:
        confidence += 0.20
    if tables:
        confidence += 0.15
    if has_pricing_schedule:
        confidence += 0.15
    if has_boq_table:
        confidence += 0.15
    if line_items:
        confidence += 0.25
        avg = sum(float(i.get("confidence") or 0.0) for i in line_items) / max(len(line_items), 1)
        confidence += min(0.10, avg * 0.10)

    return round(min(confidence, 1.0), 4)


def analyse_docx_main_document(main_document_path: str | Path) -> Dict[str, Any]:
    """
    Main entry point for DOCX MAIN DOCUMENT INTELLIGENCE ENGINE.
    """
    _ensure_dirs()

    path = Path(_safe_str(main_document_path))
    title = path.stem or "main-document"

    if not path.exists():
        result = {
            "status": "failed",
            "engine_version": ENGINE_VERSION,
            "analysed_at": _now_iso(),
            "main_document_path": str(path),
            "error": "main_document_path_not_found",
            "document_text": "",
            "paragraphs": [],
            "table_count": 0,
            "extracted_tables": [],
            "bid_reference": "",
            "bid_description": "",
            "closing_date": "",
            "briefing_required": False,
            "has_pricing_schedule": False,
            "has_boq_table": False,
            "quantity_verified": False,
            "extracted_line_items": [],
            "confidence": 0.0,
        }
        report_path = REPORT_DIR / f"{_slug(title)}__docx_main_document_intelligence_report.json"
        report_path.write_text(json.dumps(result, indent=2, default=str))
        result["report_path"] = str(report_path)
        return result

    if path.suffix.lower() != ".docx":
        result = {
            "status": "failed",
            "engine_version": ENGINE_VERSION,
            "analysed_at": _now_iso(),
            "main_document_path": str(path),
            "error": "unsupported_file_type_expected_docx",
            "document_text": "",
            "paragraphs": [],
            "table_count": 0,
            "extracted_tables": [],
            "bid_reference": "",
            "bid_description": "",
            "closing_date": "",
            "briefing_required": False,
            "has_pricing_schedule": False,
            "has_boq_table": False,
            "quantity_verified": False,
            "extracted_line_items": [],
            "confidence": 0.0,
        }
        report_path = REPORT_DIR / f"{_slug(title)}__docx_main_document_intelligence_report.json"
        report_path.write_text(json.dumps(result, indent=2, default=str))
        result["report_path"] = str(report_path)
        return result

    parsed = _read_docx(path)
    text = _safe_str(parsed.get("text"))[:MAX_TEXT_CHARS]
    paragraphs = parsed.get("paragraphs") if isinstance(parsed.get("paragraphs"), list) else []
    tables = parsed.get("tables") if isinstance(parsed.get("tables"), list) else []

    table_items = _extract_table_line_items(tables)
    explicit_items = _extract_explicit_quantity_rows(paragraphs)
    line_items = _dedupe_items(table_items + explicit_items)

    has_pricing_schedule = _has_pricing_schedule(text, tables)
    has_boq_table = _has_boq_table(text, tables)
    briefing_required = _briefing_required(text)
    quantity_verified = bool(line_items)

    result: Dict[str, Any] = {
        "status": "ok" if parsed.get("status") == "ok" else "failed",
        "engine_version": ENGINE_VERSION,
        "analysed_at": _now_iso(),
        "main_document_path": str(path),
        "error": parsed.get("error", ""),
        "document_text": text,
        "paragraphs": paragraphs,
        "table_count": len(tables),
        "extracted_tables": _serialise_tables(tables),
        "bid_reference": _find_bid_reference(text),
        "bid_description": _find_bid_description(text, paragraphs),
        "closing_date": _find_closing_date(text),
        "briefing_required": briefing_required,
        "compulsory_briefing_required": briefing_required,
        "delivery_requirements": _delivery_requirements(text, paragraphs),
        "scope_specification_sections": _scope_or_spec_sections(paragraphs),
        "has_pricing_schedule": has_pricing_schedule,
        "has_boq_table": has_boq_table,
        "quantity_verified": quantity_verified,
        "extracted_line_items": line_items,
        "confidence": _overall_confidence(text, tables, line_items, has_pricing_schedule, has_boq_table),
        "reject_hallucinated_quantities": True,
    }
    if has_boq_table and not has_pricing_schedule and re.search(r"\b(?:boq|bill of quantities)\b", text, flags=re.I):
        requirement_source_type = "standalone_boq"
    elif has_pricing_schedule:
        requirement_source_type = "embedded_pricing_schedule"
    elif scope_or_spec_sections := result.get("scope_specification_sections"):
        requirement_source_type = "specification_schedule" if scope_or_spec_sections else "main_document_table"
    else:
        requirement_source_type = "main_document_table"

    requirement_rows = normalize_requirement_rows(
        line_items,
        source_type=requirement_source_type,
        source_document=str(path),
        default_confidence=float(result.get("confidence") or 0.0),
        evidence=["rfq_docx_main_document_intelligence"],
    )
    requirement_pack = build_requirement_pack(
        requirement_rows,
        reference_number=str(result.get("bid_reference") or ""),
        title=str(result.get("bid_description") or title),
    )
    attach_requirement_pack_fields(result, requirement_pack)

    report_path = REPORT_DIR / f"{_slug(title)}__docx_main_document_intelligence_report.json"
    report_path.write_text(json.dumps(result, indent=2, default=str))
    result["report_path"] = str(report_path)

    return result


def analyse_main_docx(main_document_path: str | Path) -> Dict[str, Any]:
    """
    Alias for API compatibility.
    """
    return analyse_docx_main_document(main_document_path)


def get_status() -> Dict[str, Any]:
    return {
        "status": "ok",
        "engine_version": ENGINE_VERSION,
        "runtime_dir": str(RUNTIME_DIR),
        "report_dir": str(REPORT_DIR),
        "capabilities": [
            "read_docx_text",
            "extract_paragraphs",
            "extract_tables",
            "detect_bid_reference",
            "detect_bid_description",
            "detect_closing_date",
            "detect_briefing_or_site_meeting",
            "detect_delivery_requirements",
            "detect_scope_specification_sections",
            "detect_pricing_schedule_tables",
            "detect_boq_like_tables",
            "extract_verified_quantity_unit_line_items",
            "reject_hallucinated_quantities",
            "write_json_report",
        ],
    }
